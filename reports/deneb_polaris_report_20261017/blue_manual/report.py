"""Deneb · Polaris 투자 전략 설명서 본편 — 섞은 스타일 (사용자 2026-10-18 승인).

* 틀: UBS Global Wealth Report 2025 지면 구조(가로 A4 · 3 단 · 상단 목차선 · 큰 인용문), 강조색 파랑.
* 각 부의 첫 쪽: 토스식 큰 숫자 카드 — 수익과 낙폭을 짝으로.
* 차트와 표: Economist 식 표지 막대 + 굵은 선으로 시작하는 표 + 지면 아래 출처.

숫자는 calc_result.json · calc2_result.json · codex_independent/dividend_projection.json 만 읽는다.

    <python-pptx 가 있는 파이썬> reports/deneb_polaris_report_20261017/blue_manual/report.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build as b  # noqa: E402
from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION  # noqa: E402
from pptx.enum.dml import MSO_LINE_DASH_STYLE  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Cm, Pt  # noqa: E402

D = b.D
D2 = json.loads((HERE / "calc2_result.json").read_text(encoding="utf-8"))
DIV = json.loads((HERE.parent / "codex_independent/dividend_projection.json").read_text(encoding="utf-8"))
H2 = D2["historical"]
RS = D2["resampling"]
OUT = HERE / "Deneb_Polaris_투자전략설명서.pptx"

b.NAV = ["개요", "Deneb", "Polaris", "적금", "전망", "자료"]
b.FOOT = "Deneb · Polaris 투자 전략 설명서"
X1, X2, X3, CW, SPAN2, FULL = b.X1, b.X2, b.X3, b.CW, b.SPAN2, b.FULL
INK, MUTED, ACCENT, PANEL, NAVY, RULE = b.INK, b.MUTED, b.ACCENT, b.PANEL, b.NAVY, b.RULE
C = b.C
AM, AML = b.AM, b.AM_LABEL
R = b.run
eok, man, pct, months = b.eok, b.man, b.pct, b.months
text, line, box, para = b.text, b.line, b.box, b.para

prs = Presentation()
prs.slide_width, prs.slide_height = Cm(b.W), Cm(b.H)
BLANK = prs.slide_layouts[6]
PAGE = [0]
TOC: list[tuple[str, str, int]] = []

NAMES = {
    "Deneb": "Deneb",
    "Polaris": "Polaris",
    "QQQ 단독 적립": "QQQ 단독 적립",
    "QLD 단독 적립": "QLD 단독",
    "TQQQ 단독 적립": "TQQQ 단독",
    "S&P500 단독 적립(SPY)": "S&P500 단독",
    "QQQ 60 + SCHD 40 고정": "QQQ 60 · SCHD 40",
    "QLD 60 + SCHD 40 고정": "QLD 60 · SCHD 40",
    "TQQQ 60 + SCHD 40 고정": "TQQQ 60 · SCHD 40",
    "변동성 조절만": "변동성 조절만",
}
KEY = {
    "Deneb": "Deneb", "Polaris": "Polaris", "QQQ 단독 적립": "QQQ", "QLD 단독 적립": "QLD", "TQQQ 단독 적립": "TQQQ",
    "S&P500 단독 적립(SPY)": "SPY", "QQQ 60 + SCHD 40 고정": "QQQ60", "QLD 60 + SCHD 40 고정": "QLD60",
    "TQQQ 60 + SCHD 40 고정": "TQQQ60", "변동성 조절만": "VOL",
}
SRC_HIST = "실제 ETF 수정 종가(QQQ · QLD · TQQQ · SCHD · SPY · IEF · GLD), 자체 계산 · 비용 0.15% · 세전 · 환율 고정 · 소수 주"


# ── 페이지 · 부품 ──────────────────────────────────────────────────────────
def page(nav: int, toc: tuple[str, str] | None = None, dark_left: float = 0.0):
    sl = prs.slides.add_slide(BLANK)
    PAGE[0] += 1
    if dark_left:
        box(sl, 0, 0, dark_left, b.H, NAVY, name="판넬")
    b.frame(sl, nav, PAGE[0], on_dark_left=dark_left)
    if toc:
        TOC.append((toc[0], toc[1], PAGE[0]))
    return sl


def kicker(sl, t, x=X1, y=2.3, color=ACCENT):
    text(sl, x, y, 12, 0.4, t, size=8.5, color=color, bold=True, spacing=1.0, after=0, name="장 표시")


def head(sl, x, y, w, t, sub="", dark=False):
    """B 식: 가는 검은 선 + 왼쪽 작은 파란 막대 + 굵은 제목 + 회색 설명."""
    line(sl, x, y, x + w, y, "FFFFFF" if dark else INK, 0.75)
    box(sl, x, y, 0.6, 0.17, "7FB0F0" if dark else ACCENT, name="표지 막대")
    text(sl, x, y + 0.28, w, 0.5, t, size=10, bold=True, color="FFFFFF" if dark else INK, spacing=1.0, after=0,
         name="그림 제목")
    if sub:
        text(sl, x, y + 0.8, w, 0.5, sub, size=8, color="C9D6EA" if dark else MUTED, spacing=1.1, after=0,
             name="그림 설명")


_TABLE = b.table


def tbl(sl, x, y, widths, rows, **kw):
    t = _TABLE(sl, x, y, widths, rows, **kw)
    line(sl, x, y, x + sum(widths), y, INK, 2.0)
    return t


def source(sl, t: str, dark: bool = False) -> None:
    """dark = 왼쪽 판넬 쪽 — 출처를 흰 바탕 오른쪽 두 단에만."""
    x, w = (X2, SPAN2) if dark else (X1, FULL)
    text(sl, x, 19.32, w, 0.45, "출처 · " + t, size=7.5, color=MUTED, spacing=1.05, after=0, name="출처")


def notes(sl, t: str) -> None:
    sl.notes_slide.notes_text_frame.text = t


def rounded(sl, x, y, w, h, fill, radius=0.06, name="카드"):
    s = box(sl, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=name)
    s.adjustments[0] = radius
    return s


def card(sl, x, y, w, h, label, big, big_color, rows, note=""):
    """A 식 큰 숫자 카드 — 큰 숫자(수익) 아래에 낙폭 줄을 짝으로."""
    rounded(sl, x, y, w, h, PANEL)
    text(sl, x + 0.6, y + 0.5, w - 1.2, 0.5, label, size=10, color=MUTED, spacing=1.0, after=0)
    text(sl, x + 0.6, y + 1.1, w - 1.2, 1.3, big, size=27, color=big_color, bold=True, spacing=1.0, after=0)
    yy = y + 2.65
    for lab, val, col in rows:
        line(sl, x + 0.6, yy - 0.1, x + w - 0.6, yy - 0.1, "D5DCE6", 0.75)
        text(sl, x + 0.6, yy, w * 0.55, 0.5, lab, size=10, color=MUTED, spacing=1.0, after=0)
        text(sl, x + w * 0.45, yy, w * 0.55 - 0.6, 0.5, val, size=11, color=col, bold=True, spacing=1.0, after=0,
             align=PP_ALIGN.RIGHT)
        yy += 0.78
    if note:
        text(sl, x + 0.6, y + h - 0.85, w - 1.2, 0.6, note, size=8, color=MUTED, spacing=1.1, after=0)


def opener(nav, part, title, lead, cards, foot="", src=""):
    sl = page(nav, (part, title))
    kicker(sl, part)
    b.title(sl, X1, 2.75, 17, title, size=32)
    text(sl, X1, 4.75, 17.5, 2.2, lead, size=13, font=b.HEAD_FONT, spacing=1.3, after=0)
    n = len(cards)
    gap = 0.75
    w = (FULL - gap * (n - 1)) / n
    for i, cd in enumerate(cards):
        card(sl, X1 + i * (w + gap), 7.6, w, 7.2, *cd)
    if foot:
        text(sl, X1, 15.45, FULL, 1.6, foot, size=10, after=0)
    if src:
        source(sl, src)
    return sl


def arrow(sl, x1, y1, x2, y2, color=INK, pt=1.0):
    ln = line(sl, x1, y1, x2, y2, color, pt)
    lnx = ln.line._get_or_add_ln()
    etree.SubElement(lnx, qn("a:tailEnd"), type="triangle", w="med", len="med")
    return ln


def series_style(ch, specs):
    """선그래프 계열: (색, 두께, 점선)."""
    for s, (col, wpt, dash) in zip(ch.plots[0].series, specs, strict=True):
        s.smooth = False
        s.format.line.color.rgb = b.rgb(col)
        s.format.line.width = Pt(wpt)
        if dash:
            s.format.line.dash_style = dash
        mk = s._element.find(qn("c:marker"))
        if mk is None:
            mk = etree.SubElement(s._element, qn("c:marker"))
        for chd in list(mk):
            mk.remove(chd)
        etree.SubElement(mk, qn("c:symbol"), val="none")


def legend_row(sl, x, y, items):
    for t, col, dash, wpt in items:
        line(sl, x, y + 0.25, x + 0.9, y + 0.25, col, wpt, dash)
        text(sl, x + 1.05, y, 4.5, 0.5, t, size=8.5, spacing=1.0, after=0)
        x += 1.4 + 0.25 * len(t) + 0.5


def bar_chart(sl, kind, x, y, w, h, cats, values, keys, fmt, vmin=None, vmax=None, unit=None, name="막대",
              cat_high=False, label_size=8):
    ch = b.add_chart(sl, kind, x, y, w, h, cats, [("값", values)], name=name)
    b._gap(ch.plots[0], 45)
    b.color_points(ch.plots[0].series[0], keys)
    b._labels(ch.plots[0].series[0], fmt, XL_LABEL_POSITION.OUTSIDE_END, label_size, INK)
    va, ca = b._axes(ch, fmt.split(";")[0] if ";" in fmt else fmt, grid=True)
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    if unit:
        va.major_unit = unit
    if cat_high:
        ca.tick_label_position = XL_TICK_LABEL_POSITION.HIGH
    ch.has_legend = False
    b.fin(ch)
    return ch


def cluster(sl, kind, x, y, w, h, cats, series, colors, fmt, vmin=None, vmax=None, unit=None, name="묶음",
            legend=True, label_size=7.5, gap=50, overlap=0):
    ch = b.add_chart(sl, kind, x, y, w, h, cats, series, name=name)
    plot = ch.plots[0]
    b._gap(plot, gap, overlap)
    for s, col in zip(plot.series, colors, strict=True):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = b.rgb(col)
        s.invert_if_negative = False
        b._labels(s, fmt, XL_LABEL_POSITION.OUTSIDE_END, label_size, INK)
    va, ca = b._axes(ch, fmt.split(";")[0], grid=True)
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    if unit:
        va.major_unit = unit
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(8)
    b.fin(ch)
    return ch


def line_chart(sl, x, y, w, h, cats, series, specs, fmt="#,##0", skip=12, name="선"):
    ch = b.add_chart(sl, XL_CHART_TYPE.LINE, x, y, w, h, cats, series, name=name)
    series_style(ch, specs)
    b._axes(ch, fmt, grid=True)
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    if skip:
        b.tick_skip(ch, skip)
    ch.has_legend = False
    b.fin(ch)
    return ch


def three_cols(sl, y, items, size=9.5):
    for x, cp in zip((X1, X2, X3), items, strict=True):
        line(sl, x, y, x + CW, y, RULE, 0.5)
        text(sl, x, y + 0.2, CW, 19.2 - y - 0.3, [cp] if not isinstance(cp, list) or isinstance(cp[0], tuple) else cp,
             size=size, after=4)


# ═════════════════════════════════════════════════════════════════════════
# 표지 · 목차
# ═════════════════════════════════════════════════════════════════════════
def cover() -> None:
    sl = prs.slides.add_slide(BLANK)
    PAGE[0] += 1
    box(sl, 0, 0, b.W, b.H, NAVY, name="표지 바탕")
    box(sl, 19.2, 0, b.W - 19.2, b.H, "123A78", name="표지 오른쪽")
    text(sl, X1, 1.4, 12, 0.5, "투자 전략 설명서 · 2026년 10월", size=10, color="C9D6EA", spacing=1.0, after=0)
    text(sl, X1, 6.2, 17, 5, ["Deneb · Polaris", "투자 전략 설명서"], size=40, color="FFFFFF", font=b.HEAD_FONT,
         spacing=1.1, after=0)
    text(sl, X1, 11.4, 16, 2.4,
         "매달 같은 돈을 넣는 두 가지 방법을 처음 보는 사람도 이해할 수 있게 정리했습니다. "
         "QQQ만 사는 것, 적금에 넣는 것과 무엇이 다른지 같은 조건으로 비교합니다.",
         size=13, color="DCE6F5", font=b.HEAD_FONT, spacing=1.35, after=0)
    text(sl, X1, 17.6, 16, 1.2, ["계산 기준: 2012년 1월 ~ 2026년 8월 실제 ETF 가격 · 세전 · 비용 0.15%",
                                 "과거 계산이며 앞으로의 수익을 약속하지 않습니다."],
         size=9, color="C9D6EA", spacing=1.3, after=0)
    yy = 6.3
    for lab, big, sub in [
        ("월 10만 원 · 14년 8개월", "넣은 돈 1,760만 원", "2012년 1월 ~ 2026년 8월, 176번"),
        ("Deneb", eok(R("Deneb")["final"]), f"최대 낙폭 {pct(R('Deneb')['mdd'])}"),
        ("Polaris", eok(R("Polaris")["final"]), f"최대 낙폭 {pct(R('Polaris')['mdd'])}"),
        ("QQQ 단독 적립", eok(R("QQQ 단독 적립")["final"]), f"최대 낙폭 {pct(R('QQQ 단독 적립')['mdd'])}"),
    ]:
        text(sl, 20.6, yy, 7.6, 0.45, lab, size=9.5, color="9FC0EC", spacing=1.0, after=0)
        text(sl, 20.6, yy + 0.5, 7.6, 0.9, big, size=19, color="FFFFFF", bold=True, spacing=1.0, after=0)
        text(sl, 20.6, yy + 1.45, 7.6, 0.45, sub, size=9, color="C9D6EA", spacing=1.0, after=0)
        yy += 2.65


def toc_page() -> tuple:
    sl = page(0)
    kicker(sl, "차례와 읽는 법")
    b.title(sl, X1, 2.75, CW, "차례")
    return sl


def fill_toc(sl) -> None:
    rows = [["장", "첫 쪽"]]
    seen = set()
    for part, t, p in TOC:
        if part in seen:
            continue
        seen.add(part)
        lab = [(part, {"bold": True, "color": ACCENT})] if "부" in part.split("·")[0] else part
        rows.append([lab, str(p)])
    tbl(sl, X1, 4.6, [CW - 1.4, 1.4], rows, size=9, row_h=0.6, head_color=MUTED)
    head(sl, X2, 2.75, SPAN2, "이 설명서를 읽는 법", "처음 나오는 용어와 숫자의 기준")
    defs = [
        ("최종 평가액", "넣은 원금과 투자이익을 더한 금액입니다. 원금이 포함돼 있으니 ‘수익’이 아닙니다. 예: 월 10만 원을 176번(1,760만 원) 넣어 8,805만 원이 됐다면 최종 평가액 8,805만 원, 투자이익 7,045만 원."),
        ("14년 8개월 적립 후", "과거 비교의 모든 금액은 2012년 1월부터 2026년 8월까지 매달 첫 거래일에 넣은 뒤 2026년 8월 31일 종가로 잰 값입니다."),
        ("연복리", "매달 넣은 돈의 효과를 뺀 ‘1원당’ 성과를 1년 평균 성장률로 바꾼 값입니다. 납입액과 상관없이 같습니다."),
        ("최대 낙폭", "1원당 가치가 가장 높았던 날에서 이후 가장 낮았던 날까지 줄어든 비율입니다. −30%면 1,000만 원이 700만 원이 된 순간이 있었다는 뜻입니다. 0 아래로 내려가는 막대로 그립니다."),
        ("회복 기간", "최대 낙폭이 시작된 고점에서 그 고점을 다시 넘을 때까지 걸린 시간입니다."),
        ("장기 시나리오", "2012~2026년의 1년 단위 실제 움직임을 무작위로 이어 붙여 만든 1,000개의 20년 모의 경로입니다. ‘하위 10%’는 1,000개 중 아래에서 100번째쯤의 값이며, 미래에 그 확률로 일어난다는 뜻이 아닙니다."),
    ]
    text(sl, X2, 3.95, CW, 15, [para(h, t) for h, t in defs[:3]], size=9.5, after=8)
    text(sl, X3, 3.95, CW, 15, [para(h, t) for h, t in defs[3:]], size=9.5, after=8)
    text(sl, X2, 11.6, SPAN2, 3.0, [
        [("세 가지 숫자를 구분합니다  ", {"bold": True}),
         ("과거 계산(2012~2026 실제 가격), 장기 시나리오(재조합 모의 경로), 실제 운용 기록입니다. "
          "이 설명서의 숫자는 앞의 두 가지이며 실제 계좌 기록은 쓰지 않았습니다. "
          "실제 계좌에는 세금 · 환율 · 정수 주 · 체결 차이가 있습니다.", {})],
    ], size=9.5, after=0)


# ═════════════════════════════════════════════════════════════════════════
# 1부 개요 (1 · 2장)
# ═════════════════════════════════════════════════════════════════════════
def part_overview() -> None:
    d, p, q = R("Deneb"), R("Polaris"), R("QQQ 단독 적립")
    opener(0, "1장 · 전략 개요", "Deneb와 Polaris",
           "Deneb는 성장을 크게 노리고 흔들림을 규칙으로 줄입니다. Polaris는 주식에 국채와 금을 섞어 덜 흔들리게 모읍니다. 둘 다 매달 같은 날 같은 돈을 넣습니다.",
           [
               ("Deneb · 월 10만 원 · 14년 8개월", eok(d["final"]), ACCENT,
                [("최대 낙폭", pct(d["mdd"]), INK), ("회복 기간", months(d["recovery_days"]), INK),
                 ("연복리", pct(d["cagr"]), INK)], "TQQQ 최대 60% + SCHD, 하락 방어"),
               ("Polaris · 월 10만 원 · 14년 8개월", eok(p["final"]), C["Polaris"],
                [("최대 낙폭", pct(p["mdd"]), INK), ("회복 기간", months(p["recovery_days"]), INK),
                 ("연복리", pct(p["cagr"]), INK)], "미국 주식 60 · 국채 20 · 금 20"),
               ("QQQ 단독 적립 · 대조군", eok(q["final"]), C["QQQ"],
                [("최대 낙폭", pct(q["mdd"]), INK), ("회복 기간", months(q["recovery_days"]), INK),
                 ("연복리", pct(q["cagr"]), INK)], "나스닥100 ETF 하나만 계속 삼"),
           ],
           "세 카드 모두 원금 1,760만 원 · 세전입니다. 큰 숫자 아래의 낙폭은 그 숫자를 얻는 동안 견뎌야 했던 하락입니다. 2012~2026년은 미국 성장주가 강했던 기간이어서 결과가 크게 나왔습니다.",
           SRC_HIST)

    sl = page(0, ("1장 · 전략 개요", "두 전략 나란히"))
    kicker(sl, "1장 · 전략 개요")
    b.title(sl, X1, 2.75, 17, "두 전략 나란히")
    rows = [
        ["", "Deneb", "Polaris", "QQQ 단독 적립 (비교 기준)"],
        ["목적", "높은 성장을 노리되 큰 하락은 규칙으로 줄인다", "주식만 들 때보다 덜 흔들리며 모은다", "나스닥100을 그대로 따라간다"],
        ["사는 것", "TQQQ(나스닥100 일간 3배) · SCHD(미국 배당주)", "SPYM(S&P500) · QQQM(나스닥100) · IEF(미국 국채 7~10년) · GLDM(금)", "QQQ(나스닥100)"],
        ["비율", "TQQQ 최대 60%, 시장이 흔들리면 줄임 · 나머지 SCHD", "30 · 30 · 20 · 20 고정 목표", "100%"],
        ["매달 넣은 돈", "목표 비율에 맞게 두 ETF를 사고, 넘친 쪽은 판다", "목표보다 모자란 종목부터 산다(팔지 않음)", "전부 QQQ를 산다"],
        ["언제 파나", "TQQQ가 고점보다 15% 내리면 절반, 30%면 전부 · 매달 맞춤", "매년 12월 맞춤 때 넘친 종목만 (버튼 승인)", "팔지 않는다"],
        ["언제 다시 사나", "나스닥100 추세가 꺼졌다 다시 켜진 날 보관 현금으로 TQQQ", "별도 재매수 규칙 없음 · 12월 맞춤", "—"],
        ["주요 위험", "3배 상품의 큰 하락 · 되산 값이 판 값보다 비쌀 수 있음", "주식 · 국채 · 금이 함께 내릴 수 있음 · 상승장 수익 포기", "고점 대비 −35% 같은 하락을 그대로 받음"],
        ["실제 운영", "토스증권 · 매일 전날 종가로 하락 방어 확인", "나무증권 원화 주문 · 매달 버튼 · 정수 주", "—"],
    ]
    tbl(sl, X1, 4.75, [3.0, 8.3, 8.3, FULL - 19.6], rows, size=9, row_h=1.05, bold_first_col=True,
        fills={(i, 1): PANEL for i in range(9)})
    source(sl, "src/auto_trading/ops/deneb_guard.py · mf_paper.py · polaris.py, docs/current/CURRENT_SPEC.md (2026-10-17 기준 규칙)")
    notes(sl, "Polaris 실제 운영은 2026-11부터 100만 원 × 4회 뒤 매년 12월 맞춤. 이 설명서의 비교는 매달 계속 적립으로 계산(운영 설정 불변).")

    # 2장
    sl = page(0, ("2장 · 공통 비교 조건", "같은 조건으로 비교하기"))
    kicker(sl, "2장 · 공통 비교 조건")
    b.title(sl, X1, 2.75, CW, "같은 조건으로\n비교하기", h=2.4)
    text(sl, X1, 5.35, CW, 4,
         "전략끼리 공정하게 비교하려면 넣는 날, 넣는 돈, 비용을 똑같이 맞춰야 합니다. 아래 조건은 이 설명서의 모든 과거 비교에 똑같이 적용했습니다.",
         size=10.5, after=0)
    rows = [
        ["항목", "조건"],
        ["기간", "2012년 1월 ~ 2026년 8월 31일 (14년 8개월)"],
        ["납입", "매달 첫 거래일 · 176번"],
        ["월 납입액", "10만 · 20만 · 30만 · 50만 · 100만 원 각각 따로 계산"],
        ["가격 자료", "실제 ETF 수정 종가 — 배당은 자동 재투자"],
        ["거래 비용", "사고팔 때마다 0.15%"],
        ["세금", "넣지 않음 (세전)"],
        ["환율", "고정 (원 = 달러로 계산)"],
        ["주식 수", "소수 주 허용 (실제는 정수 주)"],
    ]
    tbl(sl, X1, 9.0, [2.4, CW - 2.4], rows, size=9, row_h=0.62, bold_first_col=True)
    head(sl, X2, 2.75, SPAN2, "세 종류의 숫자", "이 설명서는 앞의 두 가지만 씁니다")
    kinds = [
        ("과거 계산", "2012~2026년 실제 가격으로 규칙을 그대로 돌린 결과. 실제 그 기간에 그렇게 샀다면 어땠는지를 보여 줍니다. 다만 2012년 이후 자료라 닷컴 붕괴 · 2008년 금융위기는 들어 있지 않습니다."),
        ("장기 시나리오", "같은 2012~2026년 움직임을 1년 단위로 뒤섞어 만든 1,000개의 20년 경로. 결과의 폭을 보는 용도이며 미래 확률이 아닙니다."),
        ("실제 운용 기록", "실제 계좌의 체결 결과. 이 설명서에는 쓰지 않았습니다. 실제와 계산은 세금 · 환율 · 정수 주 · 체결 가격 때문에 다를 수 있습니다."),
    ]
    xs = [X2, X2 + SPAN2 / 3 + 0.1, X2 + 2 * SPAN2 / 3 + 0.2]
    for x, (h, t) in zip(xs, kinds, strict=True):
        rounded(sl, x, 4.05, SPAN2 / 3 - 0.35, 5.3, PANEL)
        text(sl, x + 0.4, 4.45, SPAN2 / 3 - 1.1, 0.6, h, size=11.5, bold=True, spacing=1.0, after=0)
        text(sl, x + 0.4, 5.2, SPAN2 / 3 - 1.1, 4.0, t, size=9, after=0)
    head(sl, X2, 10.0, SPAN2, "납입액이 달라지면", "세전 · 소수 주 · 비율 비용이라 금액이 정확히 비례합니다")
    rows = [["월 납입액"] + [AML[a] for a in AM], ["원금 (176번)"] + [eok(176 * a) for a in AM],
            ["배수 (10만 원 기준)"] + [f"× {a // 100000}" for a in AM]]
    tbl(sl, X2, 11.3, [3.5] + [(SPAN2 - 3.5) / 5] * 5, rows, size=9, row_h=0.62, bold_first_col=True)
    text(sl, X2, 13.6, SPAN2, 5.0, [
        "다섯 금액을 따로 계산해도 연복리 · 낙폭은 모두 같고, 최종 금액만 납입액에 비례했습니다(상대 오차 10억분의 1 이하). 그래서 수익률 그림은 한 번, 금액 표는 다섯 줄로 보여 줍니다.",
        "실제 계좌에서는 정수 주, 최소 주문 금액, 해외주식 양도세(연 250만 원 공제 뒤 22%), 환율 때문에 비례가 깨집니다. 금액이 클수록 세금 차이가 커집니다.",
    ], size=9.5, after=6)
    source(sl, "blue_manual/calc_spec.md (결과 전 고정) · calc.py — 다섯 금액 각각 실행, 월 20만 원 결과가 기존 52쪽 원장과 일치")

    # 2장 비교군
    sl = page(0, ("2장 · 공통 비교 조건", "비교 대상 11가지"))
    kicker(sl, "2장 · 공통 비교 조건")
    b.title(sl, X1, 2.75, 17, "비교 대상 11가지")
    rows = [
        ["이름", "사는 상품", "규칙", "비중 맞추기"],
        ["Deneb", "TQQQ · SCHD", "TQQQ 60% × (20% ÷ 나스닥100 20일 변동성), 최대 60% · 하락 방어(−15% 절반, −30% 전부 → 현금 보관 → 추세 복귀 때 되삼)", "매달 첫 거래일"],
        ["변동성 조절만", "TQQQ · SCHD", "Deneb에서 하락 방어만 뺌 — 방어의 효과를 보기 위한 대조군", "매달"],
        ["TQQQ 60 · SCHD 40", "TQQQ · SCHD", "60 : 40 고정 — 변동성 조절도 방어도 없음", "매달 60 : 40으로"],
        ["QLD 60 · SCHD 40", "QLD(2배) · SCHD", "60 : 40 고정", "매달"],
        ["QQQ 60 · SCHD 40", "QQQ · SCHD", "60 : 40 고정", "매달"],
        ["QQQ 단독 적립", "QQQ (나스닥100 1배)", "매달 같은 돈으로 QQQ만 사서 보유 · 팔지 않음", "없음"],
        ["QLD 단독", "QLD (나스닥100 일간 2배)", "QLD만 사서 보유", "없음"],
        ["TQQQ 단독", "TQQQ (나스닥100 일간 3배)", "TQQQ만 사서 보유", "없음"],
        ["S&P500 단독", "SPY (S&P500 1배)", "SPY만 사서 보유", "없음"],
        ["Polaris", "SPY · QQQ · IEF · GLD (실제는 SPYM · QQQM · IEF · GLDM)", "30 · 30 · 20 · 20 · 새 돈은 모자란 종목부터", "매년 12월 첫 거래일"],
        ["적금 3% · 4%", "은행 적금", "1년 만기 단리 → 원리금을 같은 금리 예금으로 해마다 재예치 · 세전", "—"],
    ]
    tbl(sl, X1, 4.65, [4.0, 6.6, 12.4, FULL - 23.0], rows, size=8.8, row_h=0.95, bold_first_col=True,
        fills={(1, j): PANEL for j in range(4)})
    text(sl, X1, 17.75, FULL, 1.4,
         "고정 섞기와 Deneb는 매달 비중을 맞추고, 단독 적립은 맞출 것이 없습니다. Polaris의 SPYM · QQQM · GLDM은 같은 지수 · 자산을 따르는 보수가 낮은 상품이라 계산에는 자료가 긴 SPY · QQQ · GLD를 썼습니다.",
        size=9, after=0)
    source(sl, SRC_HIST)


# ═════════════════════════════════════════════════════════════════════════
# 2부 Deneb (3 ~ 7장)
# ═════════════════════════════════════════════════════════════════════════
def part_deneb() -> None:
    d, v, f, t3 = R("Deneb"), R("변동성 조절만"), R("TQQQ 60 + SCHD 40 고정"), R("TQQQ 단독 적립")
    opener(1, "2부 · Deneb", "Deneb",
           "같은 3배 상품이라도 담는 비율과 파는 규칙에 따라 끝 금액과 견뎌야 할 하락이 크게 달라졌습니다. 월 10만 원, 14년 8개월 적립 후입니다.",
           [
               ("TQQQ 단독 적립", eok(t3["final"]), C["TQQQ"],
                [("최대 낙폭", pct(t3["mdd"]), INK), ("회복 기간", months(t3["recovery_days"]), INK)], "가장 많이 남았지만 −82%를 견뎌야 했음"),
               ("TQQQ 60 · SCHD 40 고정", eok(f["final"]), "7E5BB8",
                [("최대 낙폭", pct(f["mdd"]), INK), ("회복 기간", months(f["recovery_days"]), INK)], "섞기만 하고 줄이지 않음"),
               ("Deneb", eok(d["final"]), ACCENT,
                [("최대 낙폭", pct(d["mdd"]), INK), ("회복 기간", months(d["recovery_days"]), INK)], "변동성 조절 + 하락 방어"),
           ],
           "TQQQ 단독보다 덜 남았지만 낙폭은 절반 아래였습니다. 이 결과는 2012~2026년 한 기간의 것이고, 하락 방어가 다음 하락에서도 같은 효과를 낸다는 보장은 없습니다.",
           SRC_HIST)
    b.page_deneb_structure.__globals__["new_slide"] = lambda: _new_nav(1, ("3장 · Deneb의 구성", "TQQQ와 SCHD"))
    b.page_deneb_structure(0)
    sl = prs.slides[-1]
    kicker(sl, "3장 · Deneb의 구성")

    # 3장 2쪽: SCHD 도 함께 내린 해
    sl = page(1, ("3장 · Deneb의 구성", "SCHD의 역할과 한계"))
    kicker(sl, "3장 · Deneb의 구성")
    b.title(sl, X1, 2.75, CW, "SCHD의 역할과\n한계", h=2.4)
    b.quote(sl, X1, 5.4, CW, 2.6, "SCHD는 흔들림을 덜어 주는 쪽이지, 원금을 지켜 주는 자산이 아닙니다.")
    text(sl, X1, 8.4, CW, 10.5, [
        "TQQQ와 SCHD를 함께 들면, TQQQ가 크게 내리는 날 SCHD가 덜 내려 계좌 전체의 낙폭이 줄어듭니다. 매달 비율을 맞추면서 덜 내린 SCHD 일부를 팔아 많이 내린 TQQQ를 싸게 사는 효과도 생깁니다.",
        "하지만 SCHD도 주식입니다. 2020년 3월이나 2022년처럼 시장 전체가 내리면 함께 내립니다. 오른쪽 표처럼 TQQQ가 −79% 내린 2022년에 SCHD는 −3%였지만, 2018년에는 둘 다 내렸습니다.",
        f"2012~2026년 Deneb 안의 SCHD 몫은 계좌의 약 {min(H2['deneb_schd']['schd'][i] / H2['deneb_schd']['value'][i] for i in range(12, len(H2['deneb_schd']['value']))):.0%} ~ "
        f"{max(H2['deneb_schd']['schd'][i] / H2['deneb_schd']['value'][i] for i in range(12, len(H2['deneb_schd']['value']))):.0%}였습니다(2013년 이후 달말). 변동성이 크거나 하락 방어가 작동한 달에 SCHD 비중이 커졌습니다.",
    ], size=9.5, after=7)
    Y = D["yearly"]
    ys = ["2015", "2018", "2020", "2022", "2025"]
    head(sl, X2, 2.75, SPAN2, "TQQQ · SCHD · QQQ 연간 수익률", "실제 ETF 수정 종가 · 달러 기준 · 배당 포함")
    cluster(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, X2, 4.0, SPAN2, 6.6, [y + "년" for y in ys],
            [("TQQQ", [round(Y["TQQQ"][y], 2) for y in ys]), ("QQQ", [round(Y["QQQ"][y], 2) for y in ys]),
             ("SCHD", [round(Y["SCHD"][y], 2) for y in ys])],
            [C["TQQQ"], C["QQQ"], "8FB3E0"], '0%;"−"0%', -1.0, 1.25, 0.5, name="연간 수익률")
    ds = H2["deneb_schd"]
    mths = ds["months"]
    head(sl, X2, 11.3, SPAN2, "Deneb 안의 SCHD 비중 · 달말", "SCHD 평가액 ÷ 계좌 평가액(하락 방어로 보관한 현금 포함) · 월 10만 원")
    line_chart(sl, X2, 12.5, SPAN2, 6.6, [m[:4] for m in mths],
               [("SCHD 비중", [s / v_ for s, v_ in zip(ds["schd"], ds["value"], strict=True)])],
               [("8FB3E0", 2.0, None)], fmt="0%", name="SCHD 비중")
    source(sl, SRC_HIST)

    # 4장 비율 선택 근거 — 긴 합성 연구
    sl = page(1, ("4장 · 비율 선택 근거", "긴 기간 연구"))
    kicker(sl, "4장 · 비율 선택 근거")
    b.title(sl, X1, 2.75, CW, "긴 기간 연구")
    b.quote(sl, X1, 4.35, CW, 2.6, "3배 상품만 오래 들면 닷컴 붕괴 같은 하락에서 거의 0이 될 수 있었습니다.")
    text(sl, X1, 7.4, CW, 11, [
        para("무엇을 했나", "실제 TQQQ는 2010년에 나왔습니다. 그래서 1986~2026년 나스닥100의 하루 움직임으로 2배 · 3배 상품을 만들어(보수 · 금리 비용 반영) 40년을 돌려 봤습니다. 배당주는 미국 고배당 상위 30% 자료를 썼습니다."),
        para("찾은 것", "3배만 들면 연 15.7%로 1배(15.2%)와 비슷했지만 낙폭은 100%였습니다. 배당주를 40% 섞고 매달 맞추면 연 20.7%로 오히려 올랐습니다. 그래도 닷컴 붕괴 때 −98%였습니다."),
        para("연구 결과와 선택", "변동성 조절(M3)은 수익 · 샤프가 가장 좋았지만 미리 정한 ‘낙폭 5%p 이상 감소’ 기준을 2배에서 못 넘어 기각됐습니다. 사용자가 “수익과 최대 낙폭을 보고” M3를 골랐고, 뒤이은 연구로 하락 방어를 더했습니다. 60%는 비교 연구의 출발점에서 이어받은 상한이며 최적값으로 증명된 숫자가 아닙니다."),
    ], size=9.5, after=7)
    head(sl, X2, 2.75, SPAN2, "레버리지 섞기 — 1986 ~ 2026년 합성 자료", "매달 비율 맞춤 · 비용 0.15% · 세전 · 이 표의 숫자는 실제 ETF 결과(다음 쪽)와 섞어 비교하지 않음")
    rows = [["구성", "연 수익", "최대 낙폭", "최악 12개월"],
            ["나스닥100 1배만", "15.2%", "−81%", "−67%"],
            ["2배만", "18.9%", "−98%", "−93%"],
            ["2배 60 + 배당주 40", "18.6%", "−88%", "−74%"],
            ["3배만", "15.7%", "−100%", "−99%"],
            ["3배 80 + 배당주 20", "19.4%", "−100%", "−96%"],
            ["3배 60 + 배당주 40", "20.7%", "−98%", "−89%"],
            ["3배 50 + 배당주 50", "20.5%", "−95%", "−83%"],
            ["3배 60 + 채권 20 + 금 20", "19.6%", "−97%", "−89%"]]
    tbl(sl, X2, 4.05, [3.9, 1.6, 1.75, 1.6], rows, size=9, row_h=0.6, bold_first_col=True,
        fills={(6, j): PANEL for j in range(4)})
    rows = [["3배 60% 상한에서 줄이는 방법", "연 수익", "최대 낙폭", "판정"],
            ["M0 고정 60 : 40", "20.7%", "−98%", "기준"],
            ["M1 추세 스위치", "19.1%", "−72%", "통과"],
            ["M2 낙폭 브레이크", "19.5%", "−64%", "통과"],
            ["M3 변동성 맞춤 (Deneb 바탕)", "24.5%", "−79%", "기각*"]]
    tbl(sl, X3 + 0.2, 4.05, [3.7, 1.35, 1.6, CW - 6.85], rows, size=9, row_h=0.75, bold_first_col=True,
        fills={(4, j): PANEL for j in range(4)})
    text(sl, X3 + 0.2, 8.55, CW - 0.2, 2.3,
         "* 기각됐지만 사용자가 선택. M3 ‘변동성 맞춤’이 지금 Deneb의 TQQQ 60% × (20% ÷ 변동성) 규칙입니다. 하락 방어(연구 183 · 184)는 그 뒤에 더해졌습니다.",
         size=8.5, color=MUTED, after=0)
    head(sl, X2, 10.9, SPAN2, "배율을 올려 얻은 것과 늘어난 것", "1986~2026 합성 · 배당주 40% 섞기 기준")
    cluster(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, X2, 12.1, SPAN2, 6.9, ["1배만", "2배 60 · 40", "3배 60 · 40", "3배 + 변동성 맞춤"],
            [("연 수익", [0.152, 0.186, 0.207, 0.245]), ("최대 낙폭", [-0.81, -0.88, -0.98, -0.79])],
            [ACCENT, "9AA5B5"], '0%;"−"0%', -1.2, 0.4, 0.2, name="배율 효과")
    source(sl, "docs/research/181 · 182 (1986-02 ~ 2026-08, 합성 레버리지 + French 고배당 30%, 각 판정 1회) · 182 §7 사용자 결정")

    # 4장 2쪽: 실제 ETF 1·2·3배 — build 의 2쪽
    b.page_leverage.__globals__["new_slide"] = lambda: _new_nav(1, ("4장 · 비율 선택 근거", "실제 ETF로 본 1·2·3배"))
    b.page_leverage(0)
    sl = prs.slides[-1]
    kicker(sl, "4장 · 비율 선택 근거")

    # 5장 Deneb와 단순 적립 — 다섯 납입액 (레버리지 단독들)
    sl = page(1, ("5장 · Deneb와 단순 적립", "다섯 납입액 — 배율별"))
    kicker(sl, "5장 · Deneb와 단순 적립")
    b.title(sl, X1, 2.75, 17, "다섯 납입액 — 배율별")
    text(sl, X1, 4.4, 17, 1.4,
         "같은 돈을 QQQ · QLD · TQQQ · S&P500에 그대로 넣었을 때와 Deneb에 넣었을 때 14년 8개월 뒤 금액입니다.",
         size=10.5, after=0)
    ms = ["QQQ 단독 적립", "QLD 단독 적립", "TQQQ 단독 적립", "S&P500 단독 적립(SPY)", "TQQQ 60 + SCHD 40 고정", "Deneb"]
    head(sl, X1, 5.9, FULL, "14년 8개월 적립 후 최종 평가액", "2012-01 ~ 2026-08 · 176회 · 단위 만 원 · 세전")
    rows = [["월 납입액", "원금"] + [NAMES[m] for m in ms] + ["Deneb − QQQ"]]
    for a in AM:
        rows.append([AML[a], man(176 * a)] + [man(R(m, a)["final"]) for m in ms]
                    + ["+" + man(R("Deneb", a)["final"] - R("QQQ 단독 적립", a)["final"])])
    rows.append(["최대 낙폭", "—"] + [pct(R(m)["mdd"]) for m in ms] + ["—"])
    rows.append(["회복 기간", "—"] + [months(R(m)["recovery_days"]) for m in ms] + ["—"])
    w0 = 2.6
    tbl(sl, X1, 7.2, [w0, 2.1] + [(FULL - w0 - 2.1 - 2.6) / 6] * 6 + [2.6], rows, size=9, row_h=0.62,
        bold_first_col=True, fills={(i, 7): PANEL for i in range(8)})
    three_cols(sl, 13.4, [
        para("가장 많이 남은 쪽", "TQQQ 단독이 모든 금액에서 가장 많이 남았습니다. 월 100만 원이면 51억 원이 넘습니다. 그러나 2022년에 고점 대비 −82%까지 내려 3년 가까이 회복하지 못했습니다."),
        para("Deneb의 위치", "TQQQ 단독의 약 70%를 남기면서 최대 낙폭은 −32%로 QQQ 단독 적립(−35%)과 비슷했습니다. 같은 돈이면 QQQ 단독 적립보다 4.1배 많이 남았습니다."),
        para("주의", "이 차이의 상당 부분은 2012~2026년 나스닥100이 연 20% 가까이 오른 덕분입니다. 상승이 약한 기간에는 레버리지의 손실이 더 커질 수 있습니다(4장 긴 기간 연구)."),
    ])
    source(sl, SRC_HIST)

    # 5장: 다섯 납입액의 결과(build 3쪽)
    b.page_amounts.__globals__["new_slide"] = lambda: _new_nav(1, ("5장 · Deneb와 단순 적립", "다섯 납입액 — QQQ와 차이"))
    b.page_amounts(0)
    kicker(prs.slides[-1], "5장 · Deneb와 단순 적립")
    source(prs.slides[-1], SRC_HIST)

    # 5장: 기간별 (build 4쪽)
    b.page_periods.__globals__["new_slide"] = lambda: _new_nav(1, ("5장 · Deneb와 단순 적립", "기간별 자산 변화"))
    b.page_periods(0)
    kicker(prs.slides[-1], "5장 · Deneb와 단순 적립")
    source(prs.slides[-1], SRC_HIST)

    # 5장: 기간별 2 — 레버리지 단독
    periods_lev()
    crisis_page()

    # 6장 SCHD 배당
    dividend_pages()

    # 7장 방어와 복귀
    defense_pages()


def _new_nav(nav, toc):
    sl = prs.slides.add_slide(BLANK)
    PAGE[0] += 1
    TOC.append((toc[0], toc[1], PAGE[0]))
    sl._nav = nav
    return sl


def _patch_frame():
    """build 의 쪽 함수가 frame(sl, active, no) 를 부를 때 지금 쪽 번호와 부 목차를 쓰게."""
    orig = b.frame

    def frame(sl, active, no, on_dark_left=0.0):
        orig(sl, getattr(sl, "_nav", active), PAGE[0], on_dark_left)

    b.frame = frame
    return orig


def periods_lev() -> None:
    sl = page(1, ("5장 · Deneb와 단순 적립", "기간별 — 레버리지 단독"))
    kicker(sl, "5장 · Deneb와 단순 적립")
    b.title(sl, X1, 2.75, CW, "기간별 — 레버리지\n단독과 비교", h=2.4)
    m = D["monthly_10"]
    allm = sorted(m["Deneb"])
    text(sl, X1, 5.4, CW, 5, [
        "앞 쪽과 같은 네 구간입니다. 이번에는 Deneb를 QLD 단독 · TQQQ 단독과 비교합니다. 칸마다 세로축이 다릅니다.",
        f"2021년 말 TQQQ 단독은 {man(m['TQQQ 단독 적립']['2021-12'])}만 원이었지만 2022년 말 {man(m['TQQQ 단독 적립']['2022-12'])}만 원으로 줄었습니다. "
        f"같은 기간 Deneb는 {man(m['Deneb']['2021-12'])}만 원에서 {man(m['Deneb']['2022-12'])}만 원이었습니다.",
        "반대로 2023년 이후 강한 상승장에서는 TQQQ 단독이 다시 앞섰습니다. 하락 방어는 큰 하락을 줄이는 대신 반등 초기를 일부 놓칩니다.",
    ], size=9.5, after=7)
    items = [("Deneb", C["Deneb"], None, 2.25), ("TQQQ 단독", C["TQQQ"], MSO_LINE_DASH_STYLE.DASH, 1.75),
             ("QLD 단독", C["QLD"], None, 1.5), ("넣은 원금 누계", INK, MSO_LINE_DASH_STYLE.SQUARE_DOT, 1.0)]
    legend_row(sl, X2, 2.85, items)
    keys = ["Deneb", "TQQQ 단독 적립", "QLD 단독 적립"]
    pos = [(X2, 3.75), (X3, 3.75), (X2, 11.4), (X3, 11.4)]
    for (a, z), (x, y) in zip(D["panels"], pos, strict=True):
        ms_ = [mm for mm in allm if a <= mm <= z]
        paid = [(allm.index(mm) + 1) * 10 for mm in ms_]
        ser = [(k, [m[k][mm] / 1e4 for mm in ms_]) for k in keys] + [("원금", paid)]
        head(sl, x, y, CW, f"{a[:4]}~{z[:4]}년" + (" (8월까지)" if z.endswith("08") else ""),
             f"월말 평가액 · 만 원 · 구간 끝 원금 {paid[-1]:,}만 원")
        line_chart(sl, x, y + 1.25, CW, 6.2, [mm[:4] for mm in ms_], ser,
                   [(c_, w_, d_) for _, c_, d_, w_ in items], name=f"레버리지 {a[:4]}")
    source(sl, SRC_HIST)


def crisis_page() -> None:
    sl = page(1, ("5장 · Deneb와 단순 적립", "하락 구간과 회복"))
    kicker(sl, "5장 · Deneb와 단순 적립")
    b.title(sl, X1, 2.75, CW, "하락 구간과 회복")
    b.quote(sl, X1, 4.35, CW, 2.2, "하락 방어가 늘 낙폭을 줄인 것은 아닙니다.")
    W_ = H2["windows"]
    text(sl, X1, 6.9, CW, 7.5, [
        f"2022년처럼 하락이 길게 이어지면 방어가 효과를 냈습니다(Deneb {pct(W_['Deneb']['2022년 한 해'])} vs TQQQ 60 · SCHD 40 고정 {pct(W_['TQQQ 60 + SCHD 40 고정']['2022년 한 해'])}).",
        f"2018년 4분기 · 2020년 3월처럼 짧고 급한 하락에서는 Deneb도 {pct(W_['Deneb']['2018년 4분기'])} · {pct(W_['Deneb']['2020년 코로나'])}로 QQQ 단독 적립({pct(W_['QQQ 단독 적립']['2018년 4분기'])} · {pct(W_['QQQ 단독 적립']['2020년 코로나'])})보다 더 내렸습니다. 팔기 전에 이미 많이 내렸기 때문입니다.",
    ], size=9.5, after=7)
    wins = list(W_["Deneb"])
    ms = ["Deneb", "QQQ 단독 적립", "TQQQ 단독 적립", "TQQQ 60 + SCHD 40 고정", "S&P500 단독 적립(SPY)", "Polaris"]
    head(sl, X2, 2.75, SPAN2, "위기 구간 낙폭", "구간 안 고점 → 구간 안 저점 · 1원당 가치(납입 효과 제외) · 일간 종가")
    cluster(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, X2, 3.95, SPAN2, 7.6, wins,
            [(NAMES[k], [round(W_[k][w], 3) for w in wins]) for k in ms],
            [C[KEY[k]] for k in ms], '0%;"−"0%', -1.0, 0, 0.25, name="위기 낙폭", gap=60, label_size=7)
    head(sl, X1, 13.9, FULL, "가장 크게 줄었던 금액 · 다섯 납입액", "일간 평가액 고점 → 이후 저점 · 그사이 넣은 돈 포함 · 단위 만 원")
    ms2 = ["Deneb", "QQQ 단독 적립", "TQQQ 단독 적립", "Polaris"]
    rows = [["월 납입액"] + [NAMES[k] for k in ms2] + ["Deneb 구간"]]
    for a in AM:
        rows.append([AML[a]] + [man(R(k, a)["money_drop"]) for k in ms2]
                    + [f"{R('Deneb', a)['money_drop_from'][:7]} → {R('Deneb', a)['money_drop_to'][:7]}"])
    tbl(sl, X1, 15.2, [3.0] + [(FULL - 3.0 - 5.0) / 4] * 4 + [5.0], rows, size=8.5, row_h=0.5, bold_first_col=True)
    source(sl, SRC_HIST + " · blue_manual/calc2.py 위기 구간")


def dividend_pages() -> None:
    proj = {p["name"]: p["annual"] for p in DIV["projection"]}
    grid = DIV["dividend_grid"]

    def g(case, m, y, k):
        return next(r[k] for r in grid if r["case"] == case and r["monthly"] == m and r["year"] == y)

    sl = page(1, ("6장 · Deneb 안의 SCHD 배당", "쌓이는 SCHD와 예상 배당"))
    kicker(sl, "6장 · Deneb 안의 SCHD 배당")
    b.title(sl, X1, 2.75, CW, "쌓이는 SCHD와\n예상 배당", h=2.4)
    text(sl, X1, 5.4, CW, 13.5, [
        "SCHD만 따로 산 경우가 아닙니다. Deneb를 매달 운용하면서 실제로 쌓이는 SCHD 몫에서 나올 배당을 계산했습니다. 하락 방어 · 매달 맞춤으로 SCHD 주수가 늘고 줄어드는 것까지 반영했습니다.",
        para("미래 가정 세 가지", "‘가격 정체’ 나스닥 · SCHD 가격과 주당 배당이 그대로. ‘완만한 성장’ 나스닥 연 6%, SCHD 가격 · 배당 연 3%. ‘성장 중 급락’ 같은 성장 중 5년째에 나스닥 −35% · SCHD −20%, 주당 배당 20% 감소."),
        para("공통 가정", "SCHD 시작 가격 $35 · 배당 수익률 연 3% · 환율 1,300원 · 분기 배당을 받는 즉시 SCHD로 재투자 · 세전. 2027년부터 20년."),
        para("읽는 법", "월 환산 = 그해 배당 ÷ 12. 실제로는 분기마다 들어오므로 ‘매달 이만큼 받는다’는 뜻이 아닙니다. 배당은 이미 평가액에 들어 있어 따로 더하지 않습니다."),
    ], size=9.5, after=7)
    head(sl, X2, 2.75, SPAN2, "연간 예상 SCHD 배당 · 월 10만 원", "세 가정 · 세전 · 배당 재투자 · 2027년 = 1년째 · 단위 만 원 (월 20만 원 계산의 절반, 비례)")
    yrs = [a["year"] for a in proj["완만한 성장"]]
    line_chart(sl, X2, 4.0, SPAN2, 7.0, [f"{y}년" for y in yrs],
               [(k, [a["annual_dividend"] / 2 / 1e4 for a in proj[k]]) for k in ("가격 정체", "완만한 성장", "성장 중 급락")],
               [("8A94A6", 1.75, MSO_LINE_DASH_STYLE.SQUARE_DOT), (ACCENT, 2.25, None), (C["QQQ"], 1.75, MSO_LINE_DASH_STYLE.DASH)],
               fmt="#,##0", skip=1, name="연간 배당")
    legend_row(sl, X2, 11.15, [("가격 정체", "8A94A6", MSO_LINE_DASH_STYLE.SQUARE_DOT, 1.75), ("완만한 성장", ACCENT, None, 2.25),
                               ("성장 중 급락", C["QQQ"], MSO_LINE_DASH_STYLE.DASH, 1.75)])
    head(sl, X2, 12.0, SPAN2, "20년째 · 다섯 납입액 · 세 가정", "그해 연간 배당(세전) · 괄호는 월 환산 · 단위 만 원")
    rows = [["월 납입액", "가격 정체", "완만한 성장", "성장 중 급락"]]
    for a in AM:
        rows.append([AML[a]] + [f"{g(k, a, 20, 'annual_dividend') / 1e4:,.0f}  ({g(k, a, 20, 'monthly_equivalent') / 1e4:,.1f})"
                                for k in ("가격 정체", "완만한 성장", "성장 중 급락")])
    tbl(sl, X2, 13.3, [3.2] + [(SPAN2 - 3.2) / 3] * 3, rows, size=9, row_h=0.6, bold_first_col=True,
        fills={(i, 2): PANEL for i in range(6)})
    source(sl, "codex_independent/dividend_projection.json (결과 전 고정 5c68278) — 기존 Deneb 엔진을 가정 가격 경로에 재생, 승인형 배당 단계 제외")

    sl = page(1, ("6장 · Deneb 안의 SCHD 배당", "납입액 × 기간별 배당"))
    kicker(sl, "6장 · Deneb 안의 SCHD 배당")
    b.title(sl, X1, 2.75, 17, "납입액 × 기간별 배당")
    case = "완만한 성장"
    head(sl, X1, 4.55, FULL, "SCHD 평가액 · 그해 연간 배당 · 월 환산 (세전과 미국 원천징수 15% 뺀 값)",
         "완만한 성장 가정 · 배당 재투자 · 단위 만 원 · 세후는 원천징수만 뺀 단순 참고이며 국내 추가 과세 · 재투자 차이는 반영하지 않음")
    rows = [["월 납입액", "기간", "넣은 원금", "SCHD 주수", "SCHD 평가액", "연간 배당 (세전)", "월 환산 (세전)", "월 환산 (×0.85)", "누적 배당 (세전)"]]
    for a in AM:
        for y in (5, 10, 15, 20):
            rows.append([AML[a] if y == 5 else "", f"{y}년", man(a * 12 * y), f"{g(case, a, y, 'schd_shares'):,.0f}주",
                         man(g(case, a, y, "schd_value")), man(g(case, a, y, "annual_dividend")),
                         f"{g(case, a, y, 'monthly_equivalent') / 1e4:,.1f}", f"{g(case, a, y, 'monthly_equivalent') * 0.85 / 1e4:,.1f}",
                         man(g(case, a, y, "cumulative_dividend"))])
    fills = {(i, j): PANEL for i in range(1, 21) if ((i - 1) // 4) % 2 == 1 for j in range(9)}
    tbl(sl, X1, 5.85, [2.6, 1.4, 2.6, 2.5, 3.1, 3.4, 3.4, 3.4, FULL - 22.4], rows, size=8.5, row_h=0.6,
        bold_first_col=True, fills=fills)
    text(sl, X1, 18.65, FULL, 0.6,
         "SCHD 주수는 Deneb 규칙대로 사고판 결과라 가정마다 다릅니다. ‘성장 중 급락’에서 배당이 오히려 많은 것은 SCHD 가격이 쌀 때 더 많은 주수를 샀기 때문입니다(6장 첫 쪽 표).",
         size=8.5, color=MUTED, after=0)
    source(sl, "codex_independent/dividend_projection.json · 가정은 시장 추정 기대값이 아님 · 승인형 배당 단계(TQQQ 목표 낮추기) 제외")


def defense_pages() -> None:
    sl = page(1, ("7장 · 하락 방어와 복귀", "파는 순서와 되사는 조건"))
    kicker(sl, "7장 · 하락 방어와 복귀")
    b.title(sl, X1, 2.75, 17, "파는 순서와 되사는 조건")
    text(sl, X1, 4.4, 17.5, 1.4,
         "예시: Deneb 계좌 1,000만 원 중 TQQQ 600만 원 · SCHD 400만 원. 가격은 설명용이며 실제 주문액이 아닙니다.",
         size=10.5, after=0)
    steps = [
        ("평상시", "TQQQ 600만 · SCHD 400만", "매일 전날 종가로 TQQQ가 최근 고점보다 얼마나 내렸는지 봄", PANEL, INK),
        ("1단계 · −15%", "TQQQ 510만 → 절반 255만 원어치 매도", "판 255만 원은 현금으로 따로 보관 · 새 돈의 TQQQ 몫도 절반으로", "DCE7F7", INK),
        ("2단계 · −30%", "남은 TQQQ 전부 매도", "보관 현금이 늘어남 · 새 돈은 전부 SCHD로", "B7CDEF", INK),
        ("기다림", "나스닥100 추세 신호가 꺼질 때까지", "20 · 60 · 120일 전보다 지금 가격이 높은지 셋 중 둘 이상이면 ‘켜짐’", "8FB0E3", INK),
        ("복귀", "꺼졌던 추세가 다시 켜진 날", "보관한 현금 전부로 TQQQ를 다시 사고, 고점을 새로 잼", ACCENT, "FFFFFF"),
    ]
    n = len(steps)
    w = (FULL - 0.6 * (n - 1)) / n
    for i, (h, l1, l2, fill, fc) in enumerate(steps):
        x = X1 + i * (w + 0.6)
        rounded(sl, x, 6.3, w, 5.4, fill, radius=0.08)
        text(sl, x + 0.35, 6.65, w - 0.7, 0.6, h, size=12, bold=True, color=fc, spacing=1.0, after=0)
        text(sl, x + 0.35, 7.5, w - 0.7, 1.4, l1, size=9.5, bold=True, color=fc, after=0)
        text(sl, x + 0.35, 9.1, w - 0.7, 2.6, l2, size=8.8, color=fc, after=0)
        if i < n - 1:
            arrow(sl, x + w + 0.05, 9.0, x + w + 0.55, 9.0, INK, 1.25)
    eps = D["brake_10"]["episodes"]
    hi = sum(e["rebuy"] > e["avg_sale"] for e in eps)
    three_cols(sl, 12.5, [
        para("판 돈은 어디에", "현금으로 보관하고 매달 비율 계산에서 뺍니다. 이자는 계산에 넣지 않았습니다. 연구에서 SCHD에 두는 방식도 비교했지만 지금 규칙은 현금입니다."),
        para("왜 추세로 되사나", "가격이 바닥인지 미리 알 수 없어서, 하락이 멈추고 다시 오르기 시작했다는 신호를 기다립니다. 그래서 바닥보다 비싸게 되사는 일이 흔합니다."),
        para("2012~2026년에", f"49번 팔고 34번 되샀습니다. 되산 값이 판 평균값보다 비쌌던 경우가 {hi}번, 쌌던 경우가 {len(eps) - hi}번이었습니다. 비싸게 되산 만큼은 방어의 비용입니다."),
    ])
    source(sl, "src/auto_trading/ops/deneb_guard.py (STAGE1 0.15 · STAGE2 0.30, 기준 T) · scripts/run_brake_184.py · calc.py brake_10")

    # 방어 있음 · 없음 · 고정
    sl = page(1, ("7장 · 하락 방어와 복귀", "방어가 있을 때와 없을 때"))
    kicker(sl, "7장 · 하락 방어와 복귀")
    b.title(sl, X1, 2.75, CW, "방어가 있을 때와\n없을 때", h=2.4)
    ms = ["TQQQ 60 + SCHD 40 고정", "변동성 조절만", "Deneb", "QQQ 단독 적립"]
    rows = [["월 10만 원", "끝 금액", "최대 낙폭", "회복"]]
    for k in ms:
        x_ = R(k)
        rows.append([NAMES[k], eok(x_["final"], ""), pct(x_["mdd"]), months(x_["recovery_days"])])
    tbl(sl, X1, 5.5, [3.25, 2.1, 1.6, 1.4], rows, size=9, row_h=0.62, bold_first_col=True,
        fills={(3, j): PANEL for j in range(4)})
    text(sl, X1, 8.9, CW, 10, [
        para("변동성 조절의 몫", f"고정 60 : 40에서 변동성 조절만 더하면 끝 금액은 비슷했고({eok(R('변동성 조절만')['final'])}) 낙폭은 −61% → −52%로 줄었습니다."),
        para("하락 방어의 몫", f"여기에 방어를 더한 Deneb는 끝 금액이 {eok(R('Deneb')['final'] - R('변동성 조절만')['final'])} 늘고 낙폭은 −32%로 줄었습니다. 2022년의 긴 하락에서 TQQQ를 일찍 판 효과가 컸습니다."),
        para("겹쳐 세지 않기", "세 줄은 같은 기간 · 같은 돈입니다. 고정 섞기 → 변동성 조절 → 방어 순으로 하나씩 더한 차이이므로 각 효과를 따로 더하면 안 됩니다."),
    ], size=9.5, after=7)
    m = D["monthly_10"]
    ks = ["Deneb", "변동성 조절만", "TQQQ 60 + SCHD 40 고정"]
    items = [("Deneb", ACCENT, None, 2.25), ("변동성 조절만", C["VOL"], MSO_LINE_DASH_STYLE.DASH, 1.75),
             ("TQQQ 60 · SCHD 40 고정", C["TQQQ60"], None, 1.75)]
    for (a, z), y, ttl in [(("2019-10", "2020-12"), 2.75, "2020년 코로나 전후"), (("2021-06", "2023-12"), 10.9, "2022년 하락 전후")]:
        mm = [x for x in sorted(m["Deneb"]) if a <= x <= z]
        head(sl, X2, y, SPAN2, f"{ttl} · 월말 평가액 확대", f"{a} ~ {z} · 월 10만 원 · 만 원 · 이 구간만의 세로축")
        line_chart(sl, X2, y + 1.2, SPAN2 - 5.2, 6.4, [x[2:4] + "." + x[5:] for x in mm],
                   [(NAMES[k], [m[k][x] / 1e4 for x in mm]) for k in ks], [(c_, w_, d_) for _, c_, d_, w_ in items],
                   skip=3, name=ttl)
        legend_x = X3 + CW - 4.9
        for i, (t_, c_, d_, w_) in enumerate(items):
            line(sl, legend_x, y + 1.6 + i * 0.7, legend_x + 0.8, y + 1.6 + i * 0.7, c_, w_, d_)
            text(sl, legend_x + 0.95, y + 1.35 + i * 0.7, 4.0, 0.5, t_, size=8.5, spacing=1.0, after=0)
    source(sl, SRC_HIST)

    # 방어의 대가
    sl = page(1, ("7장 · 하락 방어와 복귀", "방어의 대가"))
    kicker(sl, "7장 · 하락 방어와 복귀")
    b.title(sl, X1, 2.75, CW, "방어의 대가")
    eps = D["brake_10"]["episodes"]
    vals = [e["rebuy"] / e["avg_sale"] - 1 for e in eps]
    b.quote(sl, X1, 4.35, CW, 2.6, "판 값보다 비싸게 되산 만큼이 하락 방어의 비용입니다.")
    text(sl, X1, 7.3, CW, 11.5, [
        f"34번의 매도 · 되사기를 하나씩 봤습니다. 오른쪽 막대가 0보다 위면 판 평균값보다 비싸게 되산 것(놓친 상승), 아래면 싸게 되산 것(피한 하락)입니다.",
        f"가장 비싸게 되산 경우는 +{max(vals):.0%}, 가장 싸게 되산 경우는 {pct(min(vals), 0)}였습니다. ±10% 안이었던 경우는 34번 중 {sum(abs(v_) <= 0.10 for v_ in vals)}번입니다.",
        para("사고팔 때 비용", "매도와 되사기마다 0.15%씩 비용이 들고, 실제 계좌에서는 이익을 내고 판 경우 양도세도 생깁니다(이 계산은 세전)."),
        para("배당 단계는 별개", "Deneb에는 SCHD 세후 배당이 월 10 · 30 · 50 · 100만 원을 넘을 때마다 TQQQ 목표를 10%p씩 낮추자고 ‘묻는’ 규칙이 있습니다. 사용자가 승인해야 바뀌며, 이 설명서의 모든 계산에는 넣지 않았습니다."),
    ], size=9.5, after=7)
    head(sl, X2, 2.75, SPAN2, "되산 값 ÷ 판 평균값 − 1 · 34번", "2012~2026 · 위 = 비싸게 되삼(놓친 상승) · 아래 = 싸게 되삼(피한 하락) · 막대 아래 숫자는 매도 시작 연 · 월")
    cats = [e["start"][2:4] + "." + e["start"][5:7] for e in eps]
    ch = b.add_chart(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, X2, 4.0, SPAN2, 14.8, cats, [("차이", vals)], name="되사기 차이")
    b._gap(ch.plots[0], 30)
    s = ch.plots[0].series[0]
    for i, v in enumerate(vals):
        pt = s.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = b.rgb(C["QQQ"] if v > 0 else ACCENT)
    ser = s._element
    for el in ser.findall(qn("c:invertIfNegative")):
        ser.remove(el)
    dpts = ser.findall(qn("c:dPt"))
    dpts[0].addprevious(etree.Element(qn("c:invertIfNegative"), val="0"))
    for dpt in dpts:
        dpt.find(qn("c:idx")).addnext(etree.Element(qn("c:invertIfNegative"), val="0"))
    b._labels(s, '"+"0%;"−"0%;0%', XL_LABEL_POSITION.OUTSIDE_END, 6.5, INK)
    va, ca = b._axes(ch, '"+"0%;"−"0%;0%', grid=True)
    va.minimum_scale, va.maximum_scale, va.major_unit = -0.6, 1.2, 0.2
    ca.tick_labels.font.size = Pt(6.5)
    ch.has_legend = False
    b.fin(ch)
    source(sl, "blue_manual/calc_result.json brake_10.episodes (월 10만 원 실행, avg_sale = 그 보관 기간 매도 평균가, rebuy = 되산 날 가격)")


# ═════════════════════════════════════════════════════════════════════════
# 3부 Polaris (8 ~ 10장)
# ═════════════════════════════════════════════════════════════════════════
def part_polaris() -> None:
    p, q, s = R("Polaris"), R("QQQ 단독 적립"), R("S&P500 단독 적립(SPY)")
    W_ = H2["windows"]
    opener(2, "3부 · Polaris", "Polaris",
           "주식 60%에 국채 20%와 금 20%를 섞으면, 하락은 줄었지만 상승장에서 덜 올랐습니다. 월 10만 원, 14년 8개월 적립 후입니다.",
           [
               ("Polaris", eok(p["final"]), C["Polaris"],
                [("최대 낙폭", pct(p["mdd"]), INK), ("2022년 한 해 낙폭", pct(W_["Polaris"]["2022년 한 해"]), INK)], "주식 60 · 국채 20 · 금 20"),
               ("S&P500 단독 적립", eok(s["final"]), C["SPY"],
                [("최대 낙폭", pct(s["mdd"]), INK), ("2022년 한 해 낙폭", pct(W_["S&P500 단독 적립(SPY)"]["2022년 한 해"]), INK)], "미국 대형주 500곳"),
               ("QQQ 단독 적립", eok(q["final"]), C["QQQ"],
                [("최대 낙폭", pct(q["mdd"]), INK), ("2022년 한 해 낙폭", pct(W_["QQQ 단독 적립"]["2022년 한 해"]), INK)], "나스닥100"),
           ],
           f"Polaris는 QQQ 단독 적립보다 {eok(q['final'] - p['final'])} 적게 남았고, 낙폭은 {pct(q['mdd'])} → {pct(p['mdd'])}로 줄었습니다. 덜 흔들린 대가로 덜 벌었습니다.",
           SRC_HIST)
    # 8장 구성 — build 5쪽 왼쪽 판넬 재사용 대신 새 쪽
    sl = page(2, ("8장 · Polaris의 자산 배분", "네 가지 자산"), dark_left=X1 + CW + 0.4)
    kicker(sl, "8장 · Polaris의 자산 배분", color="9FC0EC")
    b.title(sl, X1, 2.75, CW - 0.2, "네 가지 자산", color="FFFFFF")
    text(sl, X1, 4.4, CW - 0.3, 2.6, "미국 주식 60%에 미국 국채 20%와 금 20%를 섞어, 주식만 들고 있을 때보다 흔들림을 줄이려는 적립식 배분입니다.",
         size=10.5, color="FFFFFF", after=0)
    ch = b.add_chart(sl, XL_CHART_TYPE.DOUGHNUT, X1 - 0.2, 7.0, 4.4, 4.4, ["SPYM", "QQQM", "IEF", "GLDM"],
                     [("비중", [0.3, 0.3, 0.2, 0.2])], name="Polaris 구성")
    sr = ch.plots[0].series[0]
    for i, col in enumerate(["7FB0F0", "3D7FD9", "BFD6F5", C["GLD"]]):
        pt = sr.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = b.rgb(col)
        pt.format.line.color.rgb = b.rgb(NAVY)
        pt.format.line.width = Pt(1.5)
    ch.has_legend = False
    hole = ch._chartSpace.find(".//" + qn("c:holeSize"))
    if hole is not None:
        hole.set("val", "58")
    b.fin(ch)
    text(sl, X1 + 4.5, 7.15, CW - 4.7, 4.3, [
        [("30%  SPYM", {"bold": True})], "S&P500 미국 대형주 500곳",
        [("30%  QQQM", {"bold": True})], "나스닥100",
        [("20%  IEF", {"bold": True})], "미국 국채 7~10년",
        [("20%  GLDM", {"bold": True})], "금",
    ], size=8.5, color="FFFFFF", after=1, spacing=1.15)
    text(sl, X1, 12.0, CW - 0.3, 7.0, [
        para("왜 S&P500과 나스닥을 나누나", "나스닥100은 기술주 비중이 크고 더 크게 오르내립니다. S&P500을 함께 두면 금융 · 헬스케어 · 소비재 같은 업종이 섞여 한쪽 업종에 덜 몰립니다.", "FFFFFF"),
        para("겹치는 부분", "애플 · 마이크로소프트 · 엔비디아 같은 대형 기술주는 두 ETF에 모두 들어 있어, 주식 60% 안의 실제 기술주 비중은 30%보다 큽니다.", "FFFFFF"),
    ], size=9, color="DCE6F5", after=6)
    head(sl, X2, 2.75, SPAN2, "비율의 근거 — 같은 조건에서 비교한 배분", "2005~2026 · 월 적립 · 참고 계산(연구 판정 아님) · 나무 안정 계좌 비교 기록")
    rows = [["구성", "연 수익", "최대 낙폭", "메모"],
            ["S&P500 60 · 국채 20 · 금 20", "9.8%", "−32.4%", "S&P500 중심"],
            ["나스닥 60 · 국채 20 · 금 20", "12.6%", "−34.1%", "나스닥 중심"],
            ["S&P 30 · 나스닥 30 · 국채 20 · 금 20", "11.2%", "−32.7%", "지금 Polaris"],
            ["S&P500 2배 60 · 국채 20 · 금 20", "—", "−62%", "레버리지 섞기"],
            ["S&P500 3배 60 · 국채 20 · 금 20", "—", "−79%", "레버리지 섞기"]]
    tbl(sl, X2, 4.05, [6.2, 2.2, 2.4, SPAN2 - 10.8], rows, size=9, row_h=0.62, bold_first_col=True,
        fills={(3, j): PANEL for j in range(4)})
    text(sl, X2, 8.3, SPAN2, 3.2, [
        "30 · 30 · 20 · 20은 S&P500 중심(낙폭 작음)과 나스닥 중심(수익 큼)의 중간을 고른 것입니다. 낙폭 차이는 1~2%p로 작았고, 수익 차이는 연 1.4%p였습니다. 최적이라고 증명된 비율이 아니라, 안정적인 계좌를 원한 사용자의 선택입니다.",
    ], size=9.5, after=0)
    head(sl, X2, 11.6, SPAN2, "국채와 금이 도움이 된 해, 안 된 해", "실제 ETF 연간 수익률 · 달러 기준")
    Y = D["yearly"]
    ys = ["2013", "2018", "2020", "2022", "2025"]
    cluster(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, X2, 12.8, SPAN2, 6.2, [y + "년" for y in ys],
            [(lab, [round(Y[k][y], 2) for y in ys]) for k, lab in [("QQQ", "QQQ"), ("SPY", "SPY"), ("IEF", "IEF 국채"), ("GLD", "GLD 금")]],
            [C["QQQ"], C["SPY"], C["IEF"], C["GLD"]], '0%;"−"0%', -0.4, 0.8, 0.2, name="자산 연간", label_size=7)
    source(sl, "docs/research · reports/namu_core_compare (2005~2026 참고 계산, 커밋 adfa870 · 19c767e · b90e73f) · 실제 ETF 수정 종가", dark=True)

    # 9장 관리 규칙 — build 5쪽(상황표) 재사용
    b.page_polaris.__globals__["new_slide"] = lambda: _new_nav(2, ("9장 · 채권과 금 관리", "상황별 관리"))
    b.page_polaris(0)
    kicker(prs.slides[-1], "9장 · 채권과 금 관리", color="9FC0EC")
    source(prs.slides[-1], "src/auto_trading/ops/polaris.py (plan_buys · plan_rebalance) · run_namu_core_compare.run_policy YEAR", dark=True)

    # 9장: 비중 변화 + 12월 맞춤 예시
    sl = page(2, ("9장 · 채권과 금 관리", "비중 변화와 12월 맞춤"))
    kicker(sl, "9장 · 채권과 금 관리")
    b.title(sl, X1, 2.75, CW, "비중 변화와\n12월 맞춤", h=2.4)
    wy, wb = H2["polaris_weights"]["YEAR"], H2["polaris_weights"]["BUY"]
    text(sl, X1, 5.4, CW, 6.5, [
        f"매달 새 돈으로 모자란 종목을 사기 때문에, 계속 넣는 동안에는 비율이 크게 쏠리지 않았습니다. 12월 맞춤 없이 사기만 해도 주식 비중은 {min(wb['stock']):.0%} ~ {max(wb['stock']):.0%} 사이였습니다(2012~2026, 달말).",
        "넣는 돈에 비해 계좌가 커지면 새 돈의 힘이 약해져 쏠림이 커집니다. 앞선 연구에서 4회만 넣고 사기만 했을 때 주식이 82~90%까지 쏠렸습니다. 12월 맞춤은 그때를 위한 장치입니다.",
    ], size=9.5, after=7)
    head(sl, X1, 12.0, CW, "12월 맞춤 예 · 계좌 1,000만 원", "주식이 올라 68%가 된 경우 · 세전 · 비용 전 · 만 원")
    cluster(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, X1, 13.2, CW, 5.9, ["주식", "국채", "금"],
            [("맞춤 전", [680, 160, 160]), ("맞춤 후", [600, 200, 200])], ["B7CDEF", ACCENT], '#,##0', 0, 800, 200,
            name="12월 맞춤 예")
    head(sl, X2, 2.75, SPAN2, "주식(SPY + QQQ) 비중 · 달말", "2012~2026 · 월 10만 원 · 12월 맞춤(지금 규칙)과 사기만 비교 · 목표 60%")
    cats = [mm[:4] for mm in wy["months"]]
    line_chart(sl, X2, 4.0, SPAN2, 7.2, cats,
               [("12월 맞춤", wy["stock"]), ("사기만", wb["stock"]), ("목표", [0.6] * len(cats))],
               [(C["Polaris"], 2.25, None), (C["QQQ"], 1.75, MSO_LINE_DASH_STYLE.DASH), (INK, 1.0, MSO_LINE_DASH_STYLE.SQUARE_DOT)],
               fmt="0%", name="주식 비중").value_axis.minimum_scale = 0.4
    prs.slides[-1].shapes[-1].chart.value_axis.maximum_scale = 0.75
    legend_row(sl, X2, 11.35, [("12월 맞춤", C["Polaris"], None, 2.25), ("사기만", C["QQQ"], MSO_LINE_DASH_STYLE.DASH, 1.75),
                               ("목표 60%", INK, MSO_LINE_DASH_STYLE.SQUARE_DOT, 1.0)])
    head(sl, X2, 12.2, SPAN2, "국채 · 금 비중 · 12월 맞춤", "달말 · 각 목표 20%")
    line_chart(sl, X2, 13.4, SPAN2, 5.6, cats, [("국채", wy["bond"]), ("금", wy["gold"])],
               [(C["IEF"], 2.0, None), (C["GLD"], 2.0, MSO_LINE_DASH_STYLE.DASH)], fmt="0%", name="국채 금 비중").value_axis.minimum_scale = 0.1
    prs.slides[-1].shapes[-1].chart.value_axis.maximum_scale = 0.3
    source(sl, "blue_manual/calc2.py polaris_weights (run_policy 날마다 끝 보유 수량 × 가격) · 사기만 82~90%: 연구 커밋 fef3cf9 (4회차 뒤)")

    # 10장 Polaris vs QQQ · S&P500 · Deneb
    sl = page(2, ("10장 · Polaris와 주식 단독 적립", "다섯 납입액 비교"))
    kicker(sl, "10장 · Polaris와 주식 단독 적립")
    b.title(sl, X1, 2.75, 17, "Polaris와 주식 단독 적립")
    ms = ["Polaris", "S&P500 단독 적립(SPY)", "QQQ 단독 적립", "Deneb"]
    head(sl, X1, 4.55, FULL, "14년 8개월 적립 후 최종 평가액 · 다섯 납입액",
         "2012-01 ~ 2026-08 · 176회 · 단위 만 원 · 세전")
    rows = [["월 납입액", "원금"] + [NAMES[k] for k in ms] + ["Polaris − QQQ", "Polaris − S&P500"]]
    for a in AM:
        pv = R("Polaris", a)["final"]
        rows.append([AML[a], man(176 * a)] + [man(R(k, a)["final"]) for k in ms]
                    + [man(pv - R("QQQ 단독 적립", a)["final"]).replace("-", "−"),
                       man(pv - R("S&P500 단독 적립(SPY)", a)["final"]).replace("-", "−")])
    tbl(sl, X1, 5.85, [2.7, 2.3] + [(FULL - 5.0 - 6.4) / 4] * 4 + [3.2, 3.2], rows, size=9, row_h=0.6,
        bold_first_col=True, fills={(i, 2): PANEL for i in range(6)})
    W_ = H2["windows"]
    wins = list(W_["Polaris"])
    head(sl, X1, 10.2, CW + 4.5, "위기 구간 낙폭", "구간 안 고점 → 저점 · 1원당 가치")
    cluster(sl, XL_CHART_TYPE.COLUMN_CLUSTERED, X1, 11.4, CW + 4.5, 7.6, wins,
            [(NAMES[k], [round(W_[k][w], 3) for w in wins]) for k in ms[:3]],
            [C["Polaris"], C["SPY"], C["QQQ"]], '0%;"−"0%', -0.4, 0, 0.1, name="Polaris 위기", label_size=7)
    Yp, Yq = H2["yearly"]["Polaris"], H2["yearly"]["QQQ 단독 적립"]
    behind = [y for y in sorted(Yp) if y < "2026" and Yp[y] < Yq[y]]
    text(sl, X1 + CW + 5.2, 10.2, FULL - CW - 5.2, 9, [
        para("줄인 위험", f"네 번의 위기 구간 모두 Polaris가 가장 덜 내렸습니다. 2022년에는 {pct(W_['Polaris']['2022년 한 해'])}로 QQQ({pct(W_['QQQ 단독 적립']['2022년 한 해'])})보다 덜 내렸지만, 국채도 −15% 내려 기대만큼 막아 주지는 못했습니다."),
        para("포기한 수익", f"2012~2025년 14해 중 {len(behind)}해는 QQQ 단독 적립보다 덜 올랐습니다. 특히 2020 · 2023년처럼 주식이 크게 오른 해에 차이가 컸습니다."),
        para("회복은 비슷", f"2022년 하락에서 고점을 되찾는 데 Polaris와 QQQ 모두 {months(R('Polaris')['recovery_days'])}이 걸렸습니다. 덜 내렸지만 회복도 느렸습니다."),
    ], size=9.5, after=7)
    source(sl, SRC_HIST + " · calc2.py 위기 구간 · 연도별")


# ═════════════════════════════════════════════════════════════════════════
# 4부 적금 (11장)
# ═════════════════════════════════════════════════════════════════════════
def part_savings() -> None:
    s4, s3 = b.sav(4, 100_000), b.sav(3, 100_000)
    opener(3, "4부 · 적금과 비교", "적금과 투자 전략",
           "적금은 약속된 이자를 받고, 투자는 지나간 기간의 결과를 보여 줍니다. 같은 날 같은 돈, 14년 8개월입니다.",
           [
               ("적금 4% · 월 10만 원", eok(s4), C["SAV4"],
                [("투자이익(이자)", eok(s4 - 1760e4), INK), ("최대 낙폭", "0%", INK)], "원금과 약정 이자가 정해져 있음 · 세전"),
               ("Polaris · 월 10만 원", eok(R("Polaris")["final"]), C["Polaris"],
                [("적금 4%보다", "+" + eok(R("Polaris")["final"] - s4), INK), ("최대 낙폭", pct(R("Polaris")["mdd"]), INK)], "덜 흔들리는 투자"),
               ("Deneb · 월 10만 원", eok(R("Deneb")["final"]), ACCENT,
                [("적금 4%보다", "+" + eok(R("Deneb")["final"] - s4), INK), ("최대 낙폭", pct(R("Deneb")["mdd"]), INK)], "크게 노리는 투자"),
           ],
           f"적금 3%는 {eok(s3)}입니다. 투자 쪽 차이를 ‘더 받은 이자’로 보면 안 됩니다. 같은 기간에도 언제 끝내느냐에 따라 결과가 크게 달라지고, 원금보다 적어질 수도 있습니다.",
           "적금: 12개월 만기 단리 → 같은 금리 1년 예금 재예치, 세전 · 투자: " + SRC_HIST)
    b.page_savings.__globals__["new_slide"] = lambda: _new_nav(3, ("11장 · 적금과 비교", "다섯 납입액 · 적금 3 · 4%"))
    b.page_savings(0)
    kicker(prs.slides[-1], "11장 · 적금과 비교")
    source(prs.slides[-1], "적금: 12개월 단리 → 같은 금리 예금 재예치 · 세전, 세후 참고는 이자 15.4% 차감 · 투자: " + SRC_HIST)


# ═════════════════════════════════════════════════════════════════════════
# 5부 장기 시나리오 (12장)
# ═════════════════════════════════════════════════════════════════════════
SC = {"Deneb": "Deneb", "Polaris": "Polaris", "QQQ 단독 적립": "QQQ 단독 적립"}


def rs(name, y, k, a):
    return RS["summary"][name][str(y)][k] * a / RS["monthly_krw"]


def sav_h(r, y, a, kind="pretax"):
    return H2["savings_horizon"][f"r{r}"][str(y)][kind] * a / 100_000


def part_outlook() -> None:
    a = 100_000
    opener(4, "5부 · 장기 시나리오", "20년 뒤의 범위",
           "2012~2026년의 실제 움직임을 1년씩 뒤섞어 만든 1,000개의 20년 경로입니다. 월 10만 원, 원금 2,400만 원. 미래 예측이 아니라 결과가 얼마나 넓게 퍼질 수 있는지를 봅니다.",
           [
               ("Deneb · 20년 · 중간값", eok(rs("Deneb", 20, "p50", a)), ACCENT,
                [("하위 10%", eok(rs("Deneb", 20, "p10", a)), INK), ("모의 최저", eok(rs("Deneb", 20, "min", a)), INK)], ""),
               ("QQQ 단독 적립 · 20년 · 중간값", eok(rs("QQQ 단독 적립", 20, "p50", a)), C["QQQ"],
                [("하위 10%", eok(rs("QQQ 단독 적립", 20, "p10", a)), INK), ("모의 최저", eok(rs("QQQ 단독 적립", 20, "min", a)), INK)], ""),
               ("Polaris · 20년 · 중간값", eok(rs("Polaris", 20, "p50", a)), C["Polaris"],
                [("하위 10%", eok(rs("Polaris", 20, "p10", a)), INK), ("모의 최저", eok(rs("Polaris", 20, "min", a)), INK)], ""),
           ],
           f"같은 조건의 적금 4%는 {eok(sav_h(4, 20, a))}입니다. 재료가 미국 성장주가 강했던 2012~2026년뿐이라 숫자가 크게 나옵니다. 닷컴 붕괴나 2008년 같은 하락은 재료에 없습니다.",
           "full_report/calculation.py 재표본(씨앗 20261018 · 1,000 경로) + blue_manual/calc2.py QQQ 단독 적립 추가 · 세전 · 비례 환산")

    # 방법 + 부채꼴
    sl = page(4, ("12장 · 장기 적립 시나리오", "만드는 방법과 퍼짐"))
    kicker(sl, "12장 · 장기 적립 시나리오")
    b.title(sl, X1, 2.75, CW, "만드는 방법과\n퍼짐", h=2.4)
    text(sl, X1, 5.4, CW, 13, [
        para("1", "2012~2026년 매일의 QQQ · TQQQ · SCHD · SPY · IEF · GLD 수익률을 같은 날끼리 묶어 둡니다."),
        para("2", "그중 1년(252거래일)짜리 덩어리를 무작위로 골라 20년이 될 때까지 이어 붙입니다. 이렇게 1,000번."),
        para("3", "경로마다 Deneb · Polaris · QQQ 단독 적립을 같은 규칙 그대로(하락 방어 · 12월 맞춤 포함) 처음부터 돌립니다. 세 전략은 같은 1,000개 경로를 씁니다."),
        para("4", "5 · 10 · 15 · 20년째 금액을 줄 세워 가장 낮은 값(모의 최저), 아래에서 10% · 50%(중간) · 90% 지점을 읽습니다."),
        para("주의", "모의 최저는 1,000개 중 가장 나빴던 값일 뿐, 실제로 그보다 나빠지지 않는다는 보장선이 아닙니다. 재료 기간이 강세장이라 모든 값이 낙관적으로 기울어 있습니다."),
    ], size=9.5, after=7)
    fan = RS["fan"]
    for i, (k, col) in enumerate([("Deneb", ACCENT), ("QQQ 단독 적립", C["QQQ"]), ("Polaris", C["Polaris"])]):
        y = 2.75 + i * 5.55
        head(sl, X2, y, SPAN2, f"{k} · 월 10만 원 · 20년", "하위 10 · 중간 · 상위 10% 선과 넣은 원금 · 만 원 · 세로축은 칸마다 다름")
        mm = RS["fan_months"]
        line_chart(sl, X2, y + 1.15, SPAN2, 4.2, [f"{(m_ - 1) // 12 + 1}년" for m_ in mm],
                   [("상위 10%", [v / 2e4 for v in fan[k]["p90"]]), ("중간", [v / 2e4 for v in fan[k]["p50"]]),
                    ("하위 10%", [v / 2e4 for v in fan[k]["p10"]]), ("원금", [m_ * 10 for m_ in mm])],
                   [(col, 1.0, MSO_LINE_DASH_STYLE.DASH), (col, 2.25, None), (col, 1.0, MSO_LINE_DASH_STYLE.DASH),
                    (INK, 1.0, MSO_LINE_DASH_STYLE.SQUARE_DOT)], skip=24, name=f"부채꼴 {k}")
    source(sl, "full_report/calculation.py state_replay · blue_manual/calc2.py (Deneb · Polaris 분위가 기존 원장과 같음을 확인)")

    for yrs in (5, 10, 15, 20):
        sl = page(4, ("12장 · 장기 적립 시나리오", f"{yrs}년 적립"))
        kicker(sl, "12장 · 장기 적립 시나리오")
        b.title(sl, X1, 2.75, 12, f"{yrs}년 적립", size=27)
        text(sl, X1 + 7.5, 3.0, FULL - 7.5, 1.6,
             f"월 납입액마다 {yrs * 12}번 넣은 뒤의 금액입니다. 같은 1,000개 경로 · 세전 · 단위 만 원. 적금은 같은 기간 같은 돈의 세전 금액입니다.",
             size=10, after=0)
        y = 4.75
        for k, col in [("Deneb", ACCENT), ("QQQ 단독 적립", C["QQQ"]), ("Polaris", C["Polaris"])]:
            head(sl, X1, y, FULL, f"{k} · {yrs}년 뒤", "")
            rows = [["월 납입액", "넣은 원금", "모의 최저", "하위 10%", "중간", "상위 10%", "적금 3%", "적금 4%"]]
            for a_ in AM:
                rows.append([AML[a_], man(a_ * 12 * yrs)] + [man(rs(k, yrs, q, a_)) for q in ("min", "p10", "p50", "p90")]
                            + [man(sav_h(3, yrs, a_)), man(sav_h(4, yrs, a_))])
            tbl(sl, X1, y + 0.75, [2.9, 2.9] + [(FULL - 5.8 - 5.4) / 4] * 4 + [2.7, 2.7], rows, size=8.5, row_h=0.42,
                bold_first_col=True, fills={(i, 4): PANEL for i in range(1, 6)})
            y += 4.85
        source(sl, "같은 1,000 경로(2012~2026 1년 블록 재조합) · 세전 · 소수 주 · 비용 0.15% · 월 20만 원 계산을 납입액에 비례 환산 · 적금 12개월 단리 재예치")

    # 공통 비교표
    sl = page(4, ("12장 · 장기 적립 시나리오", "세 전략 한눈에"))
    kicker(sl, "12장 · 장기 적립 시나리오")
    b.title(sl, X1, 2.75, 17, "세 전략 한눈에 · 월 10만 원")
    head(sl, X1, 4.55, FULL, "기간별 중간값과 모의 최저", "같은 1,000 경로 · 세전 · 단위 만 원")
    rows = [["기간", "넣은 원금", "Deneb 중간", "QQQ 중간", "Polaris 중간", "적금 4%", "Deneb 최저", "QQQ 최저", "Polaris 최저"]]
    for yrs in (5, 10, 15, 20):
        rows.append([f"{yrs}년", man(a * 12 * yrs)] + [man(rs(k, yrs, "p50", a)) for k in ("Deneb", "QQQ 단독 적립", "Polaris")]
                    + [man(sav_h(4, yrs, a))] + [man(rs(k, yrs, "min", a)) for k in ("Deneb", "QQQ 단독 적립", "Polaris")])
    tbl(sl, X1, 5.85, [2.2, 2.6] + [(FULL - 4.8) / 7] * 7, rows, size=9, row_h=0.62, bold_first_col=True)
    cmp = RS["compare"]
    head(sl, X1, 9.6, FULL, "같은 경로끼리 맞대 보기 — 1,000개 중 몇 개였나", "미래 확률이 아니라 이 모의 경로에서 센 개수")
    rows = [["기간", "Deneb가 QQQ보다 적었던 경로", "Polaris가 QQQ보다 적었던 경로", "원금보다 적었던 경로 · Deneb", "· QQQ", "· Polaris"]]
    for yrs in ("5", "10", "15", "20"):
        c_ = cmp[yrs]
        rows.append([f"{yrs}년", f"{c_['deneb_below_qqq']}개", f"{c_['polaris_below_qqq']}개", f"{c_['below_paid']['Deneb']}개",
                     f"{c_['below_paid']['QQQ 단독 적립']}개", f"{c_['below_paid']['Polaris']}개"])
    tbl(sl, X1, 10.9, [2.2, 6.2, 6.2, 5.0, 3.5, FULL - 23.1], rows, size=9, row_h=0.6, bold_first_col=True)
    three_cols(sl, 14.4, [
        para("읽는 법", "5년처럼 짧으면 세 전략 모두 원금보다 적게 끝난 경로가 있었습니다(QQQ 26개, Deneb 16개, Polaris 13개). 기간이 길수록 줄었습니다."),
        para("Deneb와 QQQ", "같은 경로에서 Deneb가 QQQ보다 적게 남은 경우는 5년 65개, 20년 2개였습니다. 재료가 강세장이라 레버리지에 유리하게 기울어 있습니다."),
        para("Polaris와 QQQ", "Polaris는 거의 모든 경로에서 QQQ보다 적게 남았습니다. 대신 하위 10%와 중간값의 간격이 좁아 결과가 덜 퍼졌습니다."),
    ])
    source(sl, "blue_manual/calc2_result.json resampling.compare")


# ═════════════════════════════════════════════════════════════════════════
# 6부 한계 · 자료 (13 · 14장)
# ═════════════════════════════════════════════════════════════════════════
def part_limits() -> None:
    sl = page(5, ("13장 · 결과 해석과 한계", "세 전략의 장단점"))
    kicker(sl, "13장 · 결과 해석과 한계")
    b.title(sl, X1, 2.75, 17, "세 전략의 장단점")
    rows = [
        ["", "Deneb", "Polaris", "QQQ 단독 적립"],
        ["확인된 점 (2012~2026)", "QQQ 단독보다 4.1배 남기면서 최대 낙폭은 비슷(−32% vs −35%)", "네 위기 구간 모두 가장 덜 내림 · 최대 낙폭 −23%", "규칙이 없어 이해 · 실행이 가장 쉬움"],
        ["감수할 점", "짧고 급한 하락에선 더 내림(2018 · 2020) · 되사기 절반은 판 값보다 비쌌음", "QQQ보다 끝 금액이 42% 적음 · 국채도 함께 내린 해(2022)", "−35% 하락과 24개월 회복을 그대로 견뎌야 함"],
        ["특정 기간에 기댄 점", "나스닥100의 강한 상승에 크게 기댐 · 긴 합성 연구에선 섞기만으로 −98%", "국채 금리 상승기(2022)에 분산 효과가 약했음", "미국 기술주 비중이 큼"],
        ["실제와 다를 수 있는 점", "양도세(매도 · 되사기 잦음) · 정수 주 · 체결 가격", "원화 주문 · 정수 주 · 분배금 재투자 규칙 없음", "환율 · 세금"],
        ["연구 판정", "변동성 조절은 연구 판정 ‘기각’, 사용자 선택 · 하락 방어는 별도 연구로 채택", "참고 계산 기반 선택(판정 아님)", "비교 기준"],
    ]
    tbl(sl, X1, 4.6, [4.2, (FULL - 4.2) / 3, (FULL - 4.2) / 3, (FULL - 4.2) / 3], rows, size=9, row_h=1.25,
        bold_first_col=True, fills={(i, 1): PANEL for i in range(6)})
    text(sl, X1, 13.0, FULL, 1.0, "이 설명서는 한 전략을 고르라는 결론을 내리지 않습니다. 같은 조건에서 무엇을 얻고 무엇을 감수했는지를 나란히 보여 줄 뿐입니다.",
         size=10.5, font=b.HEAD_FONT, after=0)
    Yh = H2["yearly"]
    ys = [str(y) for y in range(2012, 2026)]
    head(sl, X1, 14.35, FULL, "해마다 수익률 · 1원당 가치", "2012~2025 · 실제 ETF · 세전 · 0 아래는 그해 손실")
    rows = [["전략"] + [y[2:] + "년" for y in ys]]
    for k in ("Deneb", "QQQ 단독 적립", "Polaris"):
        rows.append([NAMES[k]] + [f"{round(Yh[k][y] * 100) + 0:.0f}%".replace("-", "−") for y in ys])
    tbl(sl, X1, 15.6, [3.4] + [(FULL - 3.4) / 14] * 14, rows, size=8.3, row_h=0.55, bold_first_col=True)
    source(sl, SRC_HIST + " · calc2.py 연도별")

    sl = page(5, ("13장 · 결과 해석과 한계", "숫자의 한계"))
    kicker(sl, "13장 · 결과 해석과 한계")
    b.title(sl, X1, 2.75, CW, "숫자의 한계")
    b.quote(sl, X1, 4.35, CW, 2.6, "과거 결과를 미래의 수익률로 늘려 그리면 안 됩니다.")
    items = [
        ("기간의 운", "2012~2026년 나스닥100은 연 20% 가까이 올랐습니다. 이 기간 결과를 미래에 그대로 기대하면 레버리지 전략을 과대평가하게 됩니다. 시작이 2000년이나 2008년 직전이었다면 순서가 달라질 수 있습니다."),
        ("자료의 종류", "실제 ETF 결과(2012~)와 긴 합성 연구(1986~)는 자료 · 기간이 달라 같은 표에 섞지 않았습니다. 합성 자료는 실제 ETF 비용 · 괴리를 완전히 담지 못합니다."),
        ("세금", "모든 계산은 세전입니다. 해외 ETF는 매도 이익에서 연 250만 원을 뺀 뒤 22%, 배당은 미국 원천징수 15%가 있습니다. Deneb는 사고파는 횟수가 많아 세금 차이가 더 클 수 있습니다."),
        ("환율", "원 = 달러로 고정해 계산했습니다. 실제로는 환율이 오르내리며 원화 금액이 달라집니다."),
        ("실행 차이", "계산은 종가로 체결했다고 봅니다. 실제는 지정가 주문 · 정수 주 · 최소 주문 금액 · 체결 실패가 있습니다."),
        ("긴 회복", "최대 낙폭에서 회복까지 QQQ 24개월, TQQQ 단독 36개월이 걸렸습니다. 그동안 적립을 멈추지 않고 이어 가는 것 자체가 어려운 일입니다."),
    ]
    text(sl, X1, 7.4, CW, 11.5, [para(h, t) for h, t in items[:2]], size=9.5, after=8)
    text(sl, X2, 2.85, CW, 16, [para(h, t) for h, t in items[2:4]], size=9.5, after=8)
    text(sl, X3, 2.85, CW, 16, [para(h, t) for h, t in items[4:]], size=9.5, after=8)
    head(sl, X2, 9.6, SPAN2, "검증이 남은 부분", "")
    text(sl, X2, 10.5, SPAN2, 8.5, [
        "하락 방어가 2012년 이전 위기(2000 · 2008)에서도 같은 효과를 내는지는 실제 ETF로 확인할 수 없습니다(TQQQ는 2010년 출시).",
        "배당을 현금으로 받는 경로, 세후 재투자 경로는 계산하지 않았습니다.",
        "승인형 배당 단계(TQQQ 목표 낮추기)를 적용했을 때의 결과는 계산하지 않았습니다.",
        "Polaris의 실제 운영(100만 원 × 4회 뒤 12월 맞춤)과 이 설명서의 매달 계속 적립은 다릅니다.",
    ], size=9.5, after=6)
    source(sl, "calc_spec.md · calc2_spec.md · docs/research/181 · 182 · 183 · 184")

    # 14장 전체 결과표
    sl = page(5, ("14장 · 전체 결과표와 출처", "전체 결과표"))
    kicker(sl, "14장 · 전체 결과표와 출처")
    b.title(sl, X1, 2.75, 17, "전체 결과표 — 과거 계산")
    head(sl, X1, 4.55, FULL, "14년 8개월 적립 후 · 다섯 납입액 · 열한 가지 비교 대상",
         "2012-01 ~ 2026-08 · 176회 · 단위 만 원 · 세전 · 연복리 · 낙폭 · 회복은 납입액과 무관")
    order = ["Deneb", "변동성 조절만", "TQQQ 60 + SCHD 40 고정", "QLD 60 + SCHD 40 고정", "QQQ 60 + SCHD 40 고정",
             "QQQ 단독 적립", "QLD 단독 적립", "TQQQ 단독 적립", "S&P500 단독 적립(SPY)", "Polaris"]
    rows = [["비교 대상"] + [AML[a] for a in AM] + ["연복리", "최대 낙폭", "회복"]]
    for k in order:
        x_ = R(k)
        rows.append([NAMES[k]] + [man(R(k, a)["final"]) for a in AM] + [pct(x_["cagr"]), pct(x_["mdd"]), months(x_["recovery_days"])])
    for r_ in (3, 4):
        rows.append([f"적금 {r_}%"] + [man(b.sav(r_, a)) for a in AM] + ["—", "0%", "—"])
    rows.append(["넣은 원금"] + [man(176 * a) for a in AM] + ["", "", ""])
    tbl(sl, X1, 5.85, [4.4] + [2.75] * 5 + [2.3, 2.3, FULL - 4.4 - 13.75 - 4.6], rows, size=9, row_h=0.6,
        bold_first_col=True, fills={(1, j): PANEL for j in range(9)})
    source(sl, SRC_HIST + " · 적금 12개월 단리 재예치 세전")

    # 정의 · 가정 · 출처
    sl = page(5, ("14장 · 전체 결과표와 출처", "정의 · 가정 · 출처"))
    kicker(sl, "14장 · 전체 결과표와 출처")
    b.title(sl, X1, 2.75, 17, "정의 · 가정 · 출처")
    col1 = [
        para("가격 자료", "data/research/div_mix_180 (QQQ · SCHD · SPY · IEF · GLD), data/research/lev_div_181 (TQQQ · QLD) — 수정 종가, 배당 재투자."),
        para("과거 계산", "scripts/run_brake_184.simulate (Deneb · 고정 섞기 · 단독 적립), scripts/run_namu_core_compare.run_policy (Polaris). 함수 · 비용 · 신호는 바꾸지 않음."),
        para("다섯 납입액", "blue_manual/calc.py — 결과 전 조건 고정(16f398e), 금액마다 따로 실행. 월 20만 원 결과가 52쪽 원장과 일치, 비례 오차 1e-9 이하."),
    ]
    col2 = [
        para("장기 시나리오", "full_report/calculation.py 재표본(씨앗 20261018, 1,000 경로, 252일 블록) + blue_manual/calc2.py QQQ 단독 적립 추가(조건 고정 9d575bc). Deneb · Polaris 분위가 기존 원장과 같음을 확인."),
        para("SCHD 예상 배당", "codex_independent/dividend_projection.json (조건 고정 5c68278) — 세 가정 · 세전 · 재투자 · 환율 1,300원 · 시작 $35 · 수익률 3%."),
        para("긴 합성 연구", "docs/research/181 · 182 (1986-02 ~ 2026-08) · 하락 방어 183 · 184."),
    ]
    col3 = [
        para("규칙 원문", "src/auto_trading/ops/deneb_guard.py · mf_paper.py · polaris.py, docs/current/CURRENT_SPEC.md (2026-10-17)."),
        para("적금", "매달 같은 날 납입, 12개월 만기 단리, 만기금은 같은 금리 1년 예금으로 해마다 재예치(연 복리), 마지막 덜 찬 기간은 비례 단리. 세전, 세후 참고는 이자 15.4% 차감."),
        para("이 설명서가 아닌 것", "새 전략 채택 · 통계적 우위 판정 · 실거래 허가가 아닙니다. 미래 금액과 배당은 보장되지 않습니다."),
    ]
    for x, col in zip((X1, X2, X3), (col1, col2, col3), strict=True):
        line(sl, x, 4.6, x + CW, 4.6, INK, 2.0)
        text(sl, x, 4.85, CW, 14, col, size=9.2, after=8)
    source(sl, "reports/deneb_polaris_report_20261017/blue_manual/ (calc_spec · calc · calc2_spec · calc2 · build · report)")


# ═════════════════════════════════════════════════════════════════════════
def main() -> None:
    _patch_frame()
    b.table = tbl
    b.chart_head = lambda sl, x, y, w, h, sub="", color=INK: head(sl, x, y, w, h, sub, dark=(color != INK))
    b.prs = prs
    cover()
    toc_sl = toc_page()
    part_overview()
    part_deneb()
    part_polaris()
    part_savings()
    part_outlook()
    part_limits()
    fill_toc(toc_sl)
    prs.save(OUT)
    print(OUT, PAGE[0])


if __name__ == "__main__":
    main()
