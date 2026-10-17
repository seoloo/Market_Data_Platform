"""스타일 비교 3쪽 — 같은 숫자(3쪽 '다섯 납입액의 결과')를 A · B · C 세 가지로.

A 토스식 카드 · B Economist 차트 + 스위스 격자 · C 지금 UBS + ZEW 큰 숫자 띠.
숫자는 calc_result.json 만 읽는다 (build.py 도우미 재사용).
"""

from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import build as b  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Cm, Pt  # noqa: E402

OUT = HERE / "style_samples" / "Deneb_Polaris_스타일비교_ABC.pptx"
OUT.parent.mkdir(exist_ok=True)
prs = Presentation()
prs.slide_width, prs.slide_height = Cm(b.W), Cm(b.H)
BLANK = prs.slide_layouts[6]

AM = b.AM
R = b.run
PERIOD = "2012년 1월 ~ 2026년 8월 · 14년 8개월 · 176회 · 세전 · 실제 ETF 가격"


def rows_amounts():
    out = []
    for a in AM:
        d, q, p = R("Deneb", a), R("QQQ 단독 적립", a), R("Polaris", a)
        out.append((a, d, q, p))
    return out


def rounded(sl, x, y, w, h, fill, radius=0.06, name="카드"):
    s = b.box(sl, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=name)
    s.adjustments[0] = radius
    return s


# ── A. 토스식 — 큰 숫자 · 둥근 카드 · 짧은 문장 ───────────────────────────────
def style_a() -> None:
    BLUE, INK, GREY, CARD, LINE = "3182F6", "191F28", "6B7684", "F2F4F6", "E5E8EB"
    sl = prs.slides.add_slide(BLANK)
    b.box(sl, 0, 0, b.W, b.H, "FFFFFF", name="바탕")
    b.text(sl, 1.6, 1.1, 14, 0.5, "A · 토스식 카드", size=9, color=BLUE, bold=True, spacing=1.0, after=0)
    b.text(sl, 1.6, 1.75, 20, 1.3, "14년 8개월 동안 매달 넣었다면", size=26, color=INK, bold=True, spacing=1.0, after=0)
    b.text(sl, 1.6, 3.05, 24, 0.6, PERIOD, size=10.5, color=GREY, spacing=1.0, after=0)

    d, q, p = R("Deneb"), R("QQQ 단독 적립"), R("Polaris")
    s4 = b.sav(4, 100_000)
    rounded(sl, 1.6, 4.2, 12.6, 6.6, CARD)
    b.text(sl, 2.3, 4.75, 11, 0.5, "월 10만 원씩 Deneb에 넣었다면", size=11, color=GREY, spacing=1.0, after=0)
    b.text(sl, 2.3, 5.4, 11.5, 1.5, b.eok(d["final"]), size=36, color=BLUE, bold=True, spacing=1.0, after=0)
    b.text(sl, 2.3, 7.05, 11.5, 0.6, f"넣은 돈 1,760만 원 → 투자이익 {b.eok(d['profit'])}", size=11, color=INK,
           spacing=1.0, after=0)
    yy = 8.0
    for lab, v in [("QQQ만 샀다면", q["final"]), ("Polaris였다면", p["final"]), ("적금 4%였다면", s4)]:
        b.line(sl, 2.3, yy - 0.12, 13.5, yy - 0.12, LINE, 0.75)
        b.text(sl, 2.3, yy, 6, 0.5, lab, size=11, color=GREY, spacing=1.0, after=0)
        b.text(sl, 8.0, yy, 5.5, 0.5, b.eok(v), size=11.5, color=INK, bold=True, spacing=1.0, after=0,
               align=PP_ALIGN.RIGHT)
        yy += 0.85

    rounded(sl, 14.8, 4.2, 13.3, 6.6, CARD)
    b.text(sl, 15.5, 4.75, 12, 0.5, "그 대신 이만큼 내려간 적도 있어요", size=11, color=GREY, spacing=1.0, after=0)
    b.text(sl, 15.5, 5.4, 12, 1.5, b.pct(d["mdd"]), size=36, color=INK, bold=True, spacing=1.0, after=0)
    b.text(sl, 15.5, 7.05, 12, 0.6, "Deneb 최대 낙폭 · 2021년 11월 고점 → 2023년 3월", size=11, color=INK,
           spacing=1.0, after=0)
    yy = 8.0
    for lab, v, col in [("Deneb", d["mdd"], BLUE), ("QQQ 단독", q["mdd"], "8B95A1"), ("Polaris", p["mdd"], "8B95A1")]:
        b.text(sl, 15.5, yy, 2.6, 0.5, lab, size=11, color=GREY, spacing=1.0, after=0)
        full = 7.2
        bw = full * abs(v) / 0.4
        rounded(sl, 18.3 + full - bw, yy + 0.05, bw, 0.42, col, radius=0.5, name=f"낙폭 {lab}")
        b.text(sl, 18.3 + full + 0.25, yy, 2.0, 0.5, b.pct(v), size=11, color=INK, bold=True, spacing=1.0, after=0)
        yy += 0.85

    rounded(sl, 1.6, 11.4, 26.5, 7.75, CARD)
    b.text(sl, 2.3, 11.95, 20, 0.6, "납입액별로 보면 · 14년 8개월 뒤 금액 (만 원)", size=12.5, color=INK, bold=True,
           spacing=1.0, after=0)
    rows = [["월 납입액", "넣은 돈", "Deneb", "QQQ 단독 적립", "Polaris", "Deneb − QQQ", "Polaris − QQQ"]]
    for a, dd, qq, pp in rows_amounts():
        rows.append([b.AM_LABEL[a], b.man(dd["paid"]),
                     [(b.man(dd["final"]), {"bold": True, "color": BLUE})], b.man(qq["final"]), b.man(pp["final"]),
                     "+" + b.man(dd["final"] - qq["final"]), b.man(pp["final"] - qq["final"]).replace("-", "−")])
    w = (26.5 - 1.4) / 7
    tbl = b.table(sl, 2.3, 12.95, [w] * 7, rows, size=10.5, row_h=0.82, bold_first_col=True, head_color=GREY)
    for i in range(len(rows)):
        for j in range(7):
            tc = tbl.cell(i, j)._tc.tcPr
            for tag in ("a:lnT", "a:lnB"):
                el = tc.find(b.qn(tag))
                if el is not None:
                    for ch in list(el):
                        el.remove(ch)
                    el.set("w", "9525")
                    sf = b.etree.SubElement(el, b.qn("a:solidFill"))
                    b.etree.SubElement(sf, b.qn("a:srgbClr"), val=LINE if (tag == "a:lnB" or i > 0) else CARD)
    b.text(sl, 1.6, 19.6, 26.5, 0.5,
           "세금 · 정수 주 · 환율을 넣으면 실제 금액은 달라져요. 지나간 기간의 계산이고 앞으로의 약속이 아니에요.",
           size=9, color=GREY, spacing=1.0, after=0)


# ── B. Economist 차트 + 스위스 격자 ──────────────────────────────────────────
def style_b() -> None:
    EBLUE, INK, GREY, GRID = "006BA2", "121212", "595959", "D9D9D9"
    sl = prs.slides.add_slide(BLANK)
    b.box(sl, 0, 0, b.W, b.H, "FFFFFF", name="바탕")
    L, RR = 1.6, b.W - 1.6
    col = (RR - L - 11 * 0.5) / 12

    def cx(i: int) -> float:
        return L + i * (col + 0.5)

    b.box(sl, L, 1.0, 1.1, 0.32, EBLUE, name="표지 막대")
    b.line(sl, L, 1.32, RR, 1.32, INK, 0.75)
    b.text(sl, L, 1.7, 10, 0.5, "B · Economist 차트 + 스위스 격자", size=9, color=EBLUE, bold=True, spacing=1.0, after=0)
    b.text(sl, L, 2.35, 20, 1.4, "다섯 납입액, 14년 8개월의 결과", size=30, color=INK, bold=True, spacing=1.0, after=0)
    b.text(sl, L, 3.9, 17, 1.0,
           "같은 날, 같은 비용으로 금액만 다르게 넣었다. 수익률은 같고 남은 돈의 차이만 커진다.",
           size=13, color=GREY, spacing=1.2, after=0)

    # 왼쪽 4 칸: 번호 붙은 짧은 글
    notes = [
        ("1", "같은 비율", f"연복리는 금액과 무관하다. Deneb {b.pct(R('Deneb')['cagr'])}, QQQ 단독 적립 "
                        f"{b.pct(R('QQQ 단독 적립')['cagr'])}, Polaris {b.pct(R('Polaris')['cagr'])}."),
        ("2", "커지는 차이", "월 10만 원에서 2억 7,307만 원이던 Deneb와 QQQ의 차이가 월 100만 원에서는 27억 3,067만 원이 된다."),
        ("3", "같이 커지는 흔들림", f"월 100만 원 Deneb는 2024년 12월 → 2025년 4월에 {b.eok(R('Deneb', 1_000_000)['money_drop'])}이 줄었다."),
    ]
    y = 5.6
    for n, h, t in notes:
        b.line(sl, cx(0), y, cx(3) + col, y, INK, 0.75)
        b.text(sl, cx(0), y + 0.2, 0.8, 0.8, n, size=20, color=EBLUE, bold=True, spacing=1.0, after=0)
        b.text(sl, cx(0) + 1.0, y + 0.25, cx(3) + col - cx(0) - 1.0, 0.5, h, size=11, color=INK, bold=True,
               spacing=1.0, after=0)
        b.text(sl, cx(0) + 1.0, y + 0.85, cx(3) + col - cx(0) - 1.0, 1.7, t, size=9.5, color=INK, spacing=1.3, after=0)
        y += 2.55

    # 오른쪽 8 칸: Economist 식 차트 (위 가는 선 + 작은 파란 막대)
    x0, w0 = cx(4), RR - cx(4)
    b.line(sl, x0, 5.6, RR, 5.6, INK, 0.75)
    b.box(sl, x0, 5.6, 0.55, 0.16, EBLUE, name="차트 표지")
    b.text(sl, x0, 5.95, w0, 0.5, "QQQ 단독 적립보다 더(덜) 남은 금액", size=11.5, color=INK, bold=True, spacing=1.0, after=0)
    b.text(sl, x0, 6.5, w0, 0.5, "14년 8개월 적립 후, 2026년 8월 말, 만 원", size=9.5, color=GREY, spacing=1.0, after=0)
    cats = [b.AM_LABEL[a] for a in AM][::-1]
    dn = [(R("Deneb", a)["final"] - R("QQQ 단독 적립", a)["final"]) / 1e4 for a in AM][::-1]
    po = [(R("Polaris", a)["final"] - R("QQQ 단독 적립", a)["final"]) / 1e4 for a in AM][::-1]
    ch = b.add_chart(sl, XL_CHART_TYPE.BAR_CLUSTERED, x0, 7.1, w0, 6.2, cats,
                     [("Polaris", po), ("Deneb", dn)], name="B 차이")
    plot = ch.plots[0]
    b._gap(plot, 30, 0)
    for s, c0 in [(plot.series[0], "6FB6D9"), (plot.series[1], EBLUE)]:
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = b.rgb(c0)
        s.invert_if_negative = False
        b._labels(s, '"+"#,##0;"−"#,##0', XL_LABEL_POSITION.OUTSIDE_END, 8.5, INK)
    va, ca = b._axes(ch, '"+"#,##0;"−"#,##0;0', grid=True)
    va.major_gridlines.format.line.color.rgb = b.rgb(GRID)
    va.minimum_scale, va.maximum_scale, va.major_unit = -100000, 300000, 50000
    va.tick_label_position = XL_TICK_LABEL_POSITION.HIGH
    ca.format.line.fill.background()
    ca.tick_labels.font.color.rgb = b.rgb(INK)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.TOP
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(9)
    b.fin(ch)

    # 아래: 스위스식 표 (굵은 위 선)
    yt = 13.75
    b.line(sl, x0, yt, RR, yt, INK, 2.0)
    rows = [["월 납입액", "원금", "Deneb", "QQQ 단독", "Polaris", "Deneb 이익", "QQQ 이익", "Polaris 이익"]]
    for a, dd, qq, pp in rows_amounts():
        rows.append([b.AM_LABEL[a], b.man(dd["paid"]), b.man(dd["final"]), b.man(qq["final"]), b.man(pp["final"]),
                     b.man(dd["profit"]), b.man(qq["profit"]), b.man(pp["profit"])])
    w = w0 / 8
    b.table(sl, x0, yt + 0.1, [w] * 8, rows, size=9.5, row_h=0.62, bold_first_col=True, head_color=INK)
    b.text(sl, x0, 18.25, w0, 0.5, "단위 만 원 · 이익 = 최종 금액 − 원금", size=8.5, color=GREY, spacing=1.0, after=0)
    b.text(sl, L, 19.6, RR - L, 0.5,
           "출처: QQQ · TQQQ · SCHD · SPY · IEF · GLD 실제 ETF 수정 종가, 자체 계산 (비용 0.15%, 세전, 환율 고정, 소수 주)",
           size=8, color=GREY, spacing=1.0, after=0)


# ── C. 지금 UBS 틀 + ZEW 식 큰 숫자 띠 ───────────────────────────────────────
def style_c() -> None:
    sl = prs.slides.add_slide(BLANK)
    b.frame(sl, 3, 3)
    b.text(sl, b.X1, 2.3, 10, 0.4, "C · 지금 틀 + 큰 숫자 띠", size=8.5, color=b.ACCENT, bold=True, spacing=1.0, after=0)
    b.title(sl, b.X1, 2.75, b.CW, "다섯 납입액의 결과")
    b.text(sl, b.X1, 4.35, b.CW, 2.6,
           "오른쪽 띠는 월 10만 원을 14년 8개월 넣은 결과입니다. 막대 길이가 곧 금액입니다. "
           "아래 표에서 다섯 납입액을 이어서 봅니다.", size=10.5, after=0)
    d, q, p = R("Deneb"), R("QQQ 단독 적립"), R("Polaris")
    x0, w0 = b.X2, b.SPAN2
    b.chart_head(sl, x0, 2.75, w0, "월 10만 원 · 14년 8개월 적립 후 최종 평가액",
                 "2012년 1월 ~ 2026년 8월 31일 · 176회 · 세전 · 막대 길이 = 금액")
    items = [("Deneb", d["final"], b.ACCENT, "FFFFFF"), ("QQQ 단독 적립", q["final"], b.C["QQQ"], b.INK),
             ("Polaris", p["final"], b.C["Polaris"], b.INK), ("넣은 원금", d["paid"], "C9D0DA", b.INK)]
    lab_w = 3.4
    full = w0 - lab_w
    y = 4.15
    for name, v, col, fc in items:
        bw = max(full * v / d["final"], 0.25)
        b.text(sl, x0, y + 0.3, lab_w - 0.2, 0.6, name, size=11, bold=True, spacing=1.0, after=0)
        b.box(sl, x0 + lab_w, y, bw, 1.15, col, name=f"띠 {name}")
        big = b.eok(v)
        if bw > 7:
            b.text(sl, x0 + lab_w + 0.4, y + 0.1, bw - 0.6, 1.0, big, size=24, color=fc, bold=True, spacing=1.0,
                   after=0, anchor=MSO_ANCHOR.MIDDLE)
        else:
            b.text(sl, x0 + lab_w + bw + 0.3, y + 0.1, 8, 1.0, big, size=24, color=b.INK, bold=True, spacing=1.0,
                   after=0, anchor=MSO_ANCHOR.MIDDLE)
        y += 1.45

    yt = 10.15
    b.chart_head(sl, b.X1, yt, b.FULL, "14년 8개월 적립 후 납입액별 최종 평가액과 투자이익",
                 "2012-01 ~ 2026-08 매달 첫 거래일 납입 176회 · 단위 만 원 · 세전 · 투자이익 = 최종 평가액 − 원금")
    hdr1 = ["", "원금", "Deneb", None, "QQQ 단독 적립", None, "Polaris", None, "QQQ 단독 적립과 차이", None]
    hdr2 = ["월 납입액", "", "최종 평가액", "투자이익", "최종 평가액", "투자이익", "최종 평가액", "투자이익", "Deneb", "Polaris"]
    rows = [hdr1, hdr2]
    for a, dd, qq, pp in rows_amounts():
        rows.append([b.AM_LABEL[a], b.man(dd["paid"]), b.man(dd["final"]), b.man(dd["profit"]), b.man(qq["final"]),
                     b.man(qq["profit"]), b.man(pp["final"]), b.man(pp["profit"]),
                     "+" + b.man(dd["final"] - qq["final"]), b.man(pp["final"] - qq["final"]).replace("-", "−")])
    widths = [2.6, 2.3, 2.75, 2.75, 2.75, 2.75, 2.75, 2.75, 2.7, 2.7]
    s = sum(widths)
    widths = [w_ * b.FULL / s for w_ in widths]
    fills = {(i, 2): b.PANEL for i in range(2, 7)} | {(i, 3): b.PANEL for i in range(2, 7)}
    tbl = b.table(sl, b.X1, yt + 1.35, widths, rows, size=9, head_rows=2, row_h=0.6, bold_first_col=True, fills=fills)
    for c0 in (2, 4, 6, 8):
        tbl.cell(0, c0).merge(tbl.cell(0, c0 + 1))
    y3 = 16.35
    cols = [
        b.para("읽는 법", "최종 평가액은 원금이 포함된 금액입니다. 원금을 뺀 투자이익을 따로 봅니다."),
        b.para("같은 비율", "세전 · 소수 주 · 비율 비용 계산이라 연복리와 낙폭은 다섯 금액 모두 같습니다."),
        b.para("실제와 차이", "정수 주, 최소 주문, 양도세, 환율 때문에 실제 계좌는 정확히 비례하지 않습니다."),
    ]
    for x, cp in zip((b.X1, b.X2, b.X3), cols, strict=True):
        b.line(sl, x, y3, x + b.CW, y3, b.RULE, 0.5)
        b.text(sl, x, y3 + 0.2, b.CW, 3.0, [cp], size=9, after=0)


style_a()
style_b()
style_c()
prs.save(OUT)
print(OUT)
